"""
Generates Farm_Accounting_Inventory_System.xlsx

A connected, transaction-based farm accounting and inventory workbook:
Purchase -> Inventory -> Issue/Usage -> Pen/Batch -> Production -> Sale -> Profitability

Run:  python3 build_farm_workbook.py  [output.xlsx]
Requires: openpyxl
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

OUT = sys.argv[1] if len(sys.argv) > 1 else "Farm_Accounting_Inventory_System.xlsx"

# --------------------------------------------------------------------------------------
# Styling
# --------------------------------------------------------------------------------------
FONT = "Arial"
GREEN = "1F5130"
GREEN_LIGHT = "E2EFDA"
GREY = "595959"
GREY_LIGHT = "F2F2F2"
AMBER = "FFE699"
RED_LIGHT = "F8CBAD"
RED_DARK = "9C0006"

F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=16, bold=True, color=GREEN)
F_SUB = Font(name=FONT, size=10, italic=True, color=GREY)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_SECTION = Font(name=FONT, size=12, bold=True, color=GREEN)
F_KPI = Font(name=FONT, size=12, bold=True, color="000000")
F_INPUT = Font(name=FONT, size=10, color="0000FF")

FILL_IN_HDR = PatternFill("solid", fgColor=GREEN)
FILL_AUTO_HDR = PatternFill("solid", fgColor=GREY)
FILL_AUTO = PatternFill("solid", fgColor=GREY_LIGHT)
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_SECTION = PatternFill("solid", fgColor=GREEN_LIGHT)
FILL_RED = PatternFill("solid", fgColor=RED_LIGHT)
FILL_AMBER = PatternFill("solid", fgColor=AMBER)
FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NGN = '"₦"#,##0;[Red]-"₦"#,##0;"-"'
NGN2 = '"₦"#,##0.00;[Red]-"₦"#,##0.00;"-"'
INT = '#,##0;[Red]-#,##0;"-"'
DEC1 = '#,##0.0;[Red]-#,##0.0;"-"'
DEC2 = '#,##0.00;[Red]-#,##0.00;"-"'
WT = '0.000;-0.000;"-"'
PCT = '0.0%;-0.0%;"-"'
DMY = "DD/MM/YYYY"
TXT = "@"

MAX_ROW = 5000          # data validation / conditional formatting reach on log sheets
HDR_ROW = 4             # header row of every log table
FIRST = HDR_ROW + 1     # first data row

ALL = '"<>#ALL#"'       # SUMIFS criterion that matches everything (including blanks)


def tr(t, col):
    """Whole-column structured reference."""
    return f"{t}[[{col}]]"


def me(t, col):
    """This-row structured reference."""
    return f"{t}[[#This Row],[{col}]]"


def nested_if(pairs, final):
    """[(cond, result), ...] -> IF(c1,r1,IF(c2,r2,...final))"""
    out = final
    for cond, res in reversed(pairs):
        out = f"IF({cond},{res},{out})"
    return out


def bl(col, key):
    """Look up a BATCH_REGISTER column for a batch id expression."""
    return f"INDEX({tr('tBatches', col)},MATCH({key},{tr('tBatches', 'Batch ID')},0))"


wb = Workbook()
wb.remove(wb.active)
wb.calculation = CalcProperties(fullCalcOnLoad=True)

# --------------------------------------------------------------------------------------
# Sheet creation (spec order first, extras after)
# --------------------------------------------------------------------------------------
ORDER = ["README", "PEN_REGISTER", "BATCH_REGISTER", "FEED_PURCHASES", "FEED_ISSUES",
         "DRUG_PURCHASES", "DRUG_USAGE", "MORTALITY_LOG", "PRODUCTION_LOG", "SALES",
         "OTHER_EXPENSES", "CASHBOOK", "INVENTORY", "BATCH_PROFITABILITY", "DASHBOARD",
         "STOCK_MOVEMENTS", "REPORTS", "TRENDS", "CHECKS", "LISTS", "HELPER"]
WS = {name: wb.create_sheet(name) for name in ORDER}

TAB_COLORS = {"README": "7F7F7F", "DASHBOARD": GREEN, "INVENTORY": "2F75B5",
              "BATCH_PROFITABILITY": "2F75B5", "REPORTS": "2F75B5", "TRENDS": "2F75B5",
              "CHECKS": "C00000", "LISTS": "BF8F00", "PEN_REGISTER": "BF8F00",
              "BATCH_REGISTER": "BF8F00"}
for n, c in TAB_COLORS.items():
    WS[n].sheet_properties.tabColor = c


def define(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def style_cell(c, font=F_BASE, fill=None, fmt=None, align=None, border=None):
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = border


def sheet_title(ws, title, subtitle, legend=True):
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = subtitle
    ws["A2"].font = F_SUB
    if legend:
        ws["A3"] = "■ GREEN header = you enter / select      ■ GREY header = automatic (do not type)      " \
                   "Add a new record by typing in the first empty row directly under the table."
        ws["A3"].font = Font(name=FONT, size=9, color=GREY)
    ws.sheet_view.showGridLines = False


# --------------------------------------------------------------------------------------
# Generic table (log) builder
# --------------------------------------------------------------------------------------
class Col:
    def __init__(self, name, width=14, fmt=None, formula=None, dv=None, note=None):
        self.name, self.width, self.fmt, self.formula, self.dv, self.note = name, width, fmt, formula, dv, note

    @property
    def auto(self):
        return self.formula is not None


DV_CACHE = {}


def dv_list(name, strict=True):
    return ("list", name, strict)


DV_DATE = ("date",)
DV_NUM = ("num",)
DV_INT = ("int",)


def apply_dv(ws, spec, rng):
    kind = spec[0]
    if kind == "list":
        src = spec[1]
        f1 = src
        dv = DataValidation(type="list", formula1=f1, allow_blank=True)
        dv.showErrorMessage = spec[2]
        dv.errorTitle = "Not in list"
        dv.error = "Choose a value from the dropdown list (lists are maintained on the LISTS / register sheets)."
    elif kind == "date":
        dv = DataValidation(type="date", operator="between", formula1="36526", formula2="73050", allow_blank=True)
        dv.errorTitle = "Invalid date"
        dv.error = "Enter a valid date (DD/MM/YYYY) between 01/01/2000 and 31/12/2099."
        dv.showErrorMessage = True
    elif kind == "num":
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        dv.errorTitle = "Invalid number"
        dv.error = "Enter a number that is zero or greater (no text, no negative values)."
        dv.showErrorMessage = True
    elif kind == "int":
        dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        dv.errorTitle = "Invalid number"
        dv.error = "Enter a whole number that is zero or greater."
        dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def build_table(ws, tname, cols, rows, style="TableStyleLight1"):
    for j, col in enumerate(cols, start=1):
        c = ws.cell(row=HDR_ROW, column=j, value=col.name)
        c.font = F_HDR
        c.fill = FILL_AUTO_HDR if col.auto else FILL_IN_HDR
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = col.width
    ws.row_dimensions[HDR_ROW].height = 32
    for i, row in enumerate(rows):
        r = FIRST + i
        for j, col in enumerate(cols, start=1):
            c = ws.cell(row=r, column=j)
            if col.auto:
                c.value = "=" + col.formula
                c.fill = FILL_AUTO
            else:
                c.value = row.get(col.name)
            c.font = F_BASE
            if col.fmt:
                c.number_format = col.fmt
    last_row = FIRST + max(len(rows), 1) - 1
    ref = f"A{HDR_ROW}:{get_column_letter(len(cols))}{last_row}"
    t = Table(displayName=tname, ref=ref)
    t._initialise_columns()
    for j, col in enumerate(cols):
        t.tableColumns[j].name = col.name
        if col.auto:
            t.tableColumns[j].calculatedColumnFormula = TableFormula(attr_text=col.formula)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)
    # validation + number formats reach far below the table so new rows are covered
    for j, col in enumerate(cols, start=1):
        L = get_column_letter(j)
        if col.dv:
            apply_dv(ws, col.dv, f"{L}{FIRST}:{L}{MAX_ROW}")
        if col.note:
            from openpyxl.comments import Comment
            ws.cell(row=HDR_ROW, column=j).comment = Comment(col.note, "System")
    ws.freeze_panes = f"A{FIRST}"
    return t


def check_cf(ws, col_letter):
    rng = f"{col_letter}{FIRST}:{col_letter}{MAX_ROW}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND({col_letter}{FIRST}<>"",{col_letter}{FIRST}<>"OK")'],
        fill=FILL_RED, font=Font(name=FONT, size=10, bold=True, color=RED_DARK)))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'{col_letter}{FIRST}="OK"'], font=Font(name=FONT, size=10, color="006100")))


def col_letter(cols, name):
    for j, c in enumerate(cols, start=1):
        if c.name == name:
            return get_column_letter(j)
    raise KeyError(name)


def pen_batch_checks(T, pen="Pen ID", batch="Batch ID", d="Date"):
    """Standard validation chain for a record that must belong to a pen and batch."""
    p, b, dt = me(T, pen), me(T, batch), me(T, d)
    return [
        (f'{p}=""', '"Missing Pen ID"'),
        (f"COUNTIF({tr('tPens', 'Pen ID')},{p})=0", f'"Pen "&{p}&" is not in PEN_REGISTER"'),
        (f'{b}=""', '"Missing Batch ID"'),
        (f"COUNTIF({tr('tBatches', 'Batch ID')},{b})=0", f'"Batch "&{b}&" is not in BATCH_REGISTER"'),
        (f"COUNTIFS({tr('tBatches', 'Batch ID')},{b},{tr('tBatches', 'Pen ID')},{p})=0",
         f'"Batch "&{b}&" belongs to pen "&{bl("Pen ID", b)}'),
        (f"{dt}<{bl('Batch Start Date', b)}", '"Date is before the batch start date"'),
        (f"AND(N({bl('Closing Date', b)})>0,{dt}>{bl('Closing Date', b)})", '"Date is after the batch closing date"'),
    ]


def date_checks(T, d="Date"):
    dt = me(T, d)
    return [(f'{dt}=""', '"Missing date"'), (f"NOT(ISNUMBER({dt}))", '"Date is not a valid date"')]


def blank_guard(T, cols, expr):
    cond = ",".join(f'{me(T, c)}=""' for c in cols)
    return f'IF(AND({cond}),"",{expr})'


def moving_avg_cols(T, item, qty, cost, issue_t, issue_item, issue_qty, seq_name, before_name, avg_name, fmt_qty, fmt_cost):
    """Perpetual (moving) weighted-average cost, one step per purchase row.
    Avg after purchase = (stock on hand before x previous avg + purchase cost) / (stock on hand before + qty).
    Only rows ABOVE are referenced, so there is no circular reference. Purchases must be kept in date order."""
    k = f"(ROW()-ROW({T}[[#Headers],[Date]]))"
    it, dt = me(T, item), me(T, "Date")
    above = lambda col: f"INDEX({tr(T, col)},1):INDEX({tr(T, col)},{k}-1)"
    seq = f"COUNTIF(INDEX({tr(T, item)},1):{it},{it})"
    before = (f"MAX(0,IF({k}=1,0,SUMIFS({above(qty)},{above(item)},{it}))-SUMIFS({tr(issue_t, issue_qty)},"
              f"{tr(issue_t, issue_item)},{it},{tr(issue_t, 'Date')},\"<\"&{dt}))")
    prev = f"IF({me(T, seq_name)}=1,0,LOOKUP(2,1/({above(item)}={it}),{above(avg_name)}))"
    avg = (f"IFERROR(({me(T, before_name)}*{prev}+N({me(T, cost)}))/({me(T, before_name)}+N({me(T, qty)})),{prev})")
    return [
        Col(seq_name, 8, INT, formula=seq),
        Col(before_name, 10, fmt_qty, formula=before,
            note="Stock of this item still in store just before this purchase (earlier purchases - issues before this date)."),
        Col(avg_name, 11, fmt_cost, formula=avg,
            note="Moving weighted-average cost after this purchase = (stock before x previous average + this purchase "
                 "cost) / (stock before + this quantity). Used to cost every issue until the next purchase."),
    ]


def order_check(T):
    k = f"(ROW()-ROW({T}[[#Headers],[Date]]))"
    return (f"AND({k}>1,{me(T, 'Date')}<INDEX({tr(T, 'Date')},MAX(1,{k}-1)))",
            '"Out of date order - insert this row after earlier purchases"')


def latest_avg(T, item_col, item_ref, date_ref, avg_name):
    """Average cost from the latest purchase (of this item) dated on/before date_ref."""
    return (f"IFERROR(LOOKUP(2,1/(({tr(T, item_col)}={item_ref})*({tr(T, 'Date')}<={date_ref})),"
            f"{tr(T, avg_name)}),0)")


def animal_lookup(T, batch="Batch ID"):
    return f'IFERROR({bl("Animal Type", me(T, batch))},"")'


# --------------------------------------------------------------------------------------
# LISTS (master data & settings)
# --------------------------------------------------------------------------------------
ws = WS["LISTS"]
sheet_title(ws, "LISTS & SETTINGS", "Master lists that feed every dropdown. Add new items by typing directly "
            "under a list — the dropdowns grow automatically. Rename with care: existing records keep the old text.")
ws["A5"], ws["B5"] = "SETTINGS", ""
ws["A5"].font = F_SECTION
settings = [("Farm Name", "My Farm (edit me)", None, "FarmName"),
            ("Drug expiry warning (days)", 30, INT, "ExpiryWarnDays"),
            ("Currency", "₦ Naira (formats are fixed to ₦)", None, None),
            ("Costing method", "Moving weighted average", None, None)]
for i, (lab, val, fmt, nm) in enumerate(settings):
    r = 6 + i
    ws.cell(row=r, column=1, value=lab).font = F_BOLD
    c = ws.cell(row=r, column=2, value=val)
    c.font = F_INPUT if nm else F_BASE
    if nm:
        c.fill = FILL_INPUT
        define(nm, f"LISTS!$B${r}")
    if fmt:
        c.number_format = fmt
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 30

LIST_TABLES = [
    ("tFeedTypes", [("Feed Type", 16, None, DV_NUM and None), ("Default Bag Size kg", 11, DEC1, DV_NUM),
                    ("Reorder Level kg", 11, DEC1, DV_NUM)],
     [["Pre-Starter", 25, 50], ["Starter", 25, 100], ["Grower", 25, 100], ["Finisher", 25, 100],
      ["Layer Mash", 25, 150], ["Chick Mash", 25, 50]]),
    ("tDrugs", [("Drug Name", 22, None, None), ("Form/Concentration", 20, None, None),
                ("Stock Unit", 10, None, dv_list("L_Units")), ("Reorder Level", 10, DEC1, DV_NUM)],
     [["Multivitamin", "Soluble liquid", "ml", 200], ["Antibiotic A", "Water-soluble powder", "g", 100],
      ["ND Vaccine (Lasota)", "1000-dose vial", "doses", 1000], ["Anticoccidial", "Oral solution", "ml", 100],
      ["Dewormer", "Oral suspension", "ml", 50], ["Disinfectant", "Concentrate", "litres", 2]]),
    ("tUnits", [("Unit", 10, None, None)],
     [["ml"], ["litres"], ["g"], ["kg"], ["tablets"], ["doses"], ["vials"], ["bottles"], ["sachets"], ["pieces"]]),
    ("tAnimalTypes", [("Animal Type", 13, None, None)],
     [["Broiler"], ["Layer"], ["Cockerel"], ["Turkey"], ["Noiler"], ["Pig"], ["Goat"], ["Catfish"], ["Other"]]),
    ("tProducts", [("Product", 16, None, None), ("Sale Unit", 9, None, None),
                   ("Reduces Head Count", 10, None, dv_list('"Yes,No"'))],
     [["Live Birds", "birds", "Yes"], ["Dressed Birds", "birds", "Yes"], ["Spent Layers", "birds", "Yes"],
      ["Eggs", "crates", "No"], ["Manure", "bags", "No"], ["Live Animals", "heads", "Yes"],
      ["Other Product", "units", "No"]]),
    ("tExpCat", [("Expense Category", 26, None, None)],
     [["Day-old Chicks / Stock Purchase"], ["Labour"], ["Electricity"], ["Fuel / Generator"], ["Water"],
      ["Transport"], ["Repairs"], ["Maintenance"], ["Equipment"], ["Cleaning / Disinfection"],
      ["Litter / Bedding"], ["Veterinary Services"], ["Other farm expenses"]]),
    ("tPayMethods", [("Payment Method", 12, None, None)],
     [["Cash"], ["Transfer"], ["POS"], ["Credit"], ["Other"]]),
    ("tCauses", [("Mortality Cause", 20, None, None)],
     [["Unknown"], ["Suspected disease"], ["Heat stress"], ["Cold / Chilling"], ["Smothering / Crushing"],
      ["Injury"], ["Predator"], ["Culled"], ["Other"]]),
    ("tMeasTypes", [("Measurement Type", 16, None, None)],
     [["Weighing"], ["Egg Collection"], ["Other Output"]]),
    ("tDrugReasons", [("Drug Reason", 16, None, None)],
     [["Treatment"], ["Prevention"], ["Vaccination"], ["Vitamin / Supplement"], ["Deworming"], ["Disinfection"],
      ["Other"]]),
    ("tMoveTypes", [("Movement Type", 16, None, None)],
     [["Transfer In"], ["Transfer Out"], ["Additional Stock"], ["Count Correction"], ["Other"]]),
    ("tCashTypes", [("Cash Type", 24, None, None), ("Direction", 9, None, dv_list('"In,Out"'))],
     [["Opening Balance", "In"], ["Sale Receipt", "In"], ["Customer Credit Receipt", "In"],
      ["Capital Injection", "In"], ["Loan Received", "In"], ["Other Income", "In"],
      ["Feed Purchase", "Out"], ["Drug Purchase", "Out"], ["Expense Payment", "Out"],
      ["Supplier Credit Payment", "Out"], ["Loan Repayment", "Out"], ["Owner Drawings", "Out"],
      ["Other Payment", "Out"]]),
    ("tStaff", [("Staff", 14, None, None)], [["Farm Manager"], ["Attendant 1"], ["Attendant 2"], ["Vet"]]),
    ("tCustomers", [("Customer", 20, None, None)],
     [["Walk-in Customer"], ["Customer A (sample)"], ["Customer B (sample)"], ["Market Trader (sample)"]]),
    ("tSuppliers", [("Supplier", 20, None, None)],
     [["Feed Supplier A (sample)"], ["Feed Supplier B (sample)"], ["Vet Store (sample)"],
      ["Hatchery (sample)"]]),
]
col = 4
for tname, cols, data in LIST_TABLES:
    start = col
    for j, (h, w, fmt, dv) in enumerate(cols):
        c = ws.cell(row=HDR_ROW, column=col + j, value=h)
        c.font = F_HDR
        c.fill = FILL_IN_HDR
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col + j)].width = w
        L = get_column_letter(col + j)
        if dv:
            apply_dv(ws, dv, f"{L}{FIRST}:{L}400")
        for i, rowv in enumerate(data):
            cc = ws.cell(row=FIRST + i, column=col + j, value=rowv[j])
            cc.font = F_BASE
            if fmt:
                cc.number_format = fmt
    ref = f"{get_column_letter(start)}{HDR_ROW}:{get_column_letter(start + len(cols) - 1)}{FIRST + len(data) - 1}"
    t = Table(displayName=tname, ref=ref)
    t._initialise_columns()
    for j, (h, *_rest) in enumerate(cols):
        t.tableColumns[j].name = h
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    col += len(cols) + 1
ws.row_dimensions[HDR_ROW].height = 32
ws.freeze_panes = "A5"

define("L_FeedTypes", tr("tFeedTypes", "Feed Type"))
define("L_Drugs", tr("tDrugs", "Drug Name"))
define("L_Units", tr("tUnits", "Unit"))
define("L_Animals", tr("tAnimalTypes", "Animal Type"))
define("L_Products", tr("tProducts", "Product"))
define("L_ExpCat", tr("tExpCat", "Expense Category"))
define("L_Pay", tr("tPayMethods", "Payment Method"))
define("L_Causes", tr("tCauses", "Mortality Cause"))
define("L_Meas", tr("tMeasTypes", "Measurement Type"))
define("L_Reasons", tr("tDrugReasons", "Drug Reason"))
define("L_MoveTypes", tr("tMoveTypes", "Movement Type"))
define("L_CashTypes", tr("tCashTypes", "Cash Type"))
define("L_Staff", tr("tStaff", "Staff"))
define("L_Customers", tr("tCustomers", "Customer"))
define("L_Suppliers", tr("tSuppliers", "Supplier"))
define("L_Pens", tr("tPens", "Pen ID"))
define("L_Batches", tr("tBatches", "Batch ID"))

# --------------------------------------------------------------------------------------
# HELPER (hidden): dynamic dropdown lists + dashboard filter logic
# --------------------------------------------------------------------------------------
hs = WS["HELPER"]
hs["A1"], hs["B1"], hs["C1"], hs["D1"] = "PenChoices", "PenFilter", "BatchFilter", "AnimalFilter"
hs["A2"], hs["B2"], hs["C2"], hs["D2"] = "GENERAL", "All Pens", "All Batches", "All Types"
for i in range(1, 401):
    r = i + 2
    hs[f"A{r}"] = f'=IFERROR(INDEX({tr("tPens", "Pen ID")},{i})&"","")'
    hs[f"B{r}"] = f"=A{r}"
for i in range(1, 1001):
    hs[f"C{i + 2}"] = f'=IFERROR(INDEX({tr("tBatches", "Batch ID")},{i})&"","")'
for i in range(1, 61):
    hs[f"D{i + 2}"] = f'=IFERROR(INDEX({tr("tAnimalTypes", "Animal Type")},{i})&"","")'
define("L_PensGen", f"OFFSET(HELPER!$A$2,0,0,COUNTA({tr('tPens', 'Pen ID')})+1,1)")
define("L_PenFilter", f"OFFSET(HELPER!$B$2,0,0,COUNTA({tr('tPens', 'Pen ID')})+1,1)")
define("L_BatchFilter", f"OFFSET(HELPER!$C$2,0,0,COUNTA({tr('tBatches', 'Batch ID')})+1,1)")
define("L_AnimalFilter", f"OFFSET(HELPER!$D$2,0,0,COUNTA({tr('tAnimalTypes', 'Animal Type')})+1,1)")
hs.sheet_state = "hidden"

# --------------------------------------------------------------------------------------
# SAMPLE DATA
# --------------------------------------------------------------------------------------
S = "SAMPLE DATA - delete or replace"
D = date

pens = [
    dict(**{"Pen ID": "P001", "Pen Name": "Broiler House 1", "Farm/Location": "Main Farm", "Animal Type": "Broiler"}),
    dict(**{"Pen ID": "P002", "Pen Name": "Broiler House 2", "Farm/Location": "Main Farm", "Animal Type": "Broiler"}),
    dict(**{"Pen ID": "P003", "Pen Name": "Broiler House 3", "Farm/Location": "Main Farm", "Animal Type": "Broiler"}),
    dict(**{"Pen ID": "P004", "Pen Name": "Broiler House 4", "Farm/Location": "Main Farm", "Animal Type": "Broiler"}),
    dict(**{"Pen ID": "P005", "Pen Name": "Layer House 1", "Farm/Location": "Main Farm", "Animal Type": "Layer"}),
    dict(**{"Pen ID": "P006", "Pen Name": "Broiler House 5", "Farm/Location": "Annex", "Animal Type": "Broiler"}),
]
for p in pens:
    p["Pen Closed"] = "No"
    p["Notes"] = S

batches = [
    ("BROILER-SEP-001", "P001", "Broiler", D(2026, 9, 1), 500, "Hatchery (sample)", "Cobb 500", None, "Active"),
    ("BROILER-SEP-002", "P002", "Broiler", D(2026, 9, 1), 600, "Hatchery (sample)", "Ross 308", None, "Active"),
    ("BROILER-SEP-003", "P003", "Broiler", D(2026, 9, 2), 450, "Hatchery (sample)", "Cobb 500", None, "Active"),
    ("BROILER-SEP-004", "P004", "Broiler", D(2026, 9, 2), 550, "Hatchery (sample)", "Arbor Acres", None, "Active"),
    ("LAYER-JUN-001", "P005", "Layer", D(2026, 6, 1), 400, "Hatchery (sample)", "Isa Brown (point-of-lay)", None,
     "Active"),
    ("BROILER-JUL-001", "P006", "Broiler", D(2026, 7, 1), 300, "Hatchery (sample)", "Cobb 500", D(2026, 8, 18),
     "Closed"),
]
batch_rows = [dict(zip(["Batch ID", "Pen ID", "Animal Type", "Batch Start Date", "Initial Quantity",
                        "Source/Supplier", "Breed/Type", "Closing Date", "Status"], b), Notes=S) for b in batches]
batch_rows[-1]["Notes"] = S + ". Closed batch: history kept, pen P006 now Empty and can take a new batch."

feed_pur = [
    (D(2026, 7, 1), "Starter", "Feed Supplier A (sample)", 12, 25, 14000, "Transfer", "INV-1001"),
    (D(2026, 7, 1), "Layer Mash", "Feed Supplier A (sample)", 20, 25, 13500, "Transfer", "INV-1002"),
    (D(2026, 7, 20), "Finisher", "Feed Supplier B (sample)", 12, 25, 14500, "Cash", "RCPT-220"),
    (D(2026, 8, 30), "Starter", "Feed Supplier A (sample)", 20, 25, 15500, "Transfer", "INV-1088"),
    (D(2026, 9, 15), "Starter", "Feed Supplier B (sample)", 10, 25, 16500, "POS", "RCPT-301"),
    (D(2026, 9, 15), "Grower", "Feed Supplier A (sample)", 20, 25, 16000, "Credit", "INV-1102"),
]
feed_pur_rows = [dict(zip(["Date", "Feed Type", "Supplier", "Quantity in Bags", "Bag Size in kg", "Price per Bag",
                           "Payment Method", "Reference"], f), Notes=S) for f in feed_pur]
feed_pur_rows[3]["Notes"] = S + ". 20 bags x 25 kg = 500 kg, recorded ONCE then issued to several pens."

feed_iss = [
    (D(2026, 7, 2), "Starter", "P006", "BROILER-JUL-001", 150), (D(2026, 7, 20), "Starter", "P006", "BROILER-JUL-001", 150),
    (D(2026, 7, 25), "Finisher", "P006", "BROILER-JUL-001", 150), (D(2026, 8, 10), "Finisher", "P006", "BROILER-JUL-001", 140),
    (D(2026, 7, 2), "Layer Mash", "P005", "LAYER-JUN-001", 200), (D(2026, 8, 1), "Layer Mash", "P005", "LAYER-JUN-001", 200),
    (D(2026, 9, 1), "Layer Mash", "P005", "LAYER-JUN-001", 60),
    (D(2026, 9, 2), "Starter", "P001", "BROILER-SEP-001", 40), (D(2026, 9, 2), "Starter", "P002", "BROILER-SEP-002", 60),
    (D(2026, 9, 2), "Starter", "P003", "BROILER-SEP-003", 35), (D(2026, 9, 2), "Starter", "P004", "BROILER-SEP-004", 50),
    (D(2026, 9, 10), "Starter", "P001", "BROILER-SEP-001", 60), (D(2026, 9, 10), "Starter", "P002", "BROILER-SEP-002", 70),
    (D(2026, 9, 10), "Starter", "P003", "BROILER-SEP-003", 55), (D(2026, 9, 10), "Starter", "P004", "BROILER-SEP-004", 65),
    (D(2026, 9, 16), "Starter", "P001", "BROILER-SEP-001", 60), (D(2026, 9, 16), "Starter", "P002", "BROILER-SEP-002", 70),
    (D(2026, 9, 18), "Grower", "P003", "BROILER-SEP-003", 60), (D(2026, 9, 18), "Grower", "P004", "BROILER-SEP-004", 70),
    (D(2026, 9, 22), "Grower", "P001", "BROILER-SEP-001", 80), (D(2026, 9, 22), "Grower", "P002", "BROILER-SEP-002", 90),
    (D(2026, 9, 22), "Grower", "P003", "BROILER-SEP-003", 70), (D(2026, 9, 22), "Grower", "P004", "BROILER-SEP-004", 75),
    (D(2026, 9, 25), "Starter", "P003", "BROILER-SEP-003", 50),
]
feed_iss_rows = [dict(zip(["Date", "Feed Type", "Pen ID", "Batch ID", "Quantity Issued in kg"], f),
                      **{"Issued By": "Attendant 1", "Notes": S}) for f in feed_iss]
for k in range(7, 11):
    feed_iss_rows[k]["Notes"] = S + ". Same 500 kg Starter purchase shared across pens."

drug_pur = [
    (D(2026, 6, 1), "Dewormer", "Vet Store (sample)", 200, 15, D(2027, 5, 31), "DW0601", "Cash"),
    (D(2026, 7, 1), "Anticoccidial", "Vet Store (sample)", 500, 12, D(2026, 9, 1), "AC0701", "Cash"),
    (D(2026, 8, 28), "Multivitamin", "Vet Store (sample)", 1000, 8, D(2027, 6, 30), "MV2408", "Transfer"),
    (D(2026, 8, 28), "Antibiotic A", "Vet Store (sample)", 500, 30, D(2026, 10, 15), "AB0826", "Transfer"),
    (D(2026, 8, 28), "ND Vaccine (Lasota)", "Vet Store (sample)", 4000, 2.5, D(2026, 12, 31), "ND1001", "Transfer"),
]
drug_pur_rows = [dict(zip(["Date", "Drug Name", "Supplier", "Quantity Purchased", "Unit Cost", "Expiry Date",
                           "Batch Number", "Payment Method"], d), Notes=S) for d in drug_pur]

drug_use = [
    (D(2026, 7, 5), "Anticoccidial", "P006", "BROILER-JUL-001", "Prevention", "1 ml per litre, 3 days", 150, "AC0701"),
    (D(2026, 8, 1), "Dewormer", "P005", "LAYER-JUN-001", "Deworming", "As per label", 100, "DW0601"),
    (D(2026, 9, 3), "Multivitamin", "P001", "BROILER-SEP-001", "Vitamin / Supplement", "10 ml per 20 litres", 50, "MV2408"),
    (D(2026, 9, 3), "Multivitamin", "P002", "BROILER-SEP-002", "Vitamin / Supplement", "5 ml per 20 litres", 30, "MV2408"),
    (D(2026, 9, 3), "Multivitamin", "P003", "BROILER-SEP-003", "Vitamin / Supplement", "15 ml per 20 litres", 70, "MV2408"),
    (D(2026, 9, 8), "ND Vaccine (Lasota)", "P001", "BROILER-SEP-001", "Vaccination", "1 dose per bird", 500, "ND1001"),
    (D(2026, 9, 8), "ND Vaccine (Lasota)", "P002", "BROILER-SEP-002", "Vaccination", "1 dose per bird", 600, "ND1001"),
    (D(2026, 9, 8), "ND Vaccine (Lasota)", "P003", "BROILER-SEP-003", "Vaccination", "1 dose per bird", 450, "ND1001"),
    (D(2026, 9, 8), "ND Vaccine (Lasota)", "P004", "BROILER-SEP-004", "Vaccination", "1 dose per bird", 550, "ND1001"),
    (D(2026, 9, 12), "Antibiotic A", "P002", "BROILER-SEP-002", "Treatment", "1 g per 2 litres, 5 days", 40, "AB0826"),
]
drug_use_rows = [dict(zip(["Date", "Drug Name", "Pen ID", "Batch ID", "Reason", "Dosage", "Quantity Used",
                           "Drug Batch Number"], d), **{"Administered By": "Farm Manager", "Notes": S}) for d in drug_use]
for k in (2, 3, 4):
    drug_use_rows[k]["Notes"] = S + ". Same drug, different pens & dosages = separate records."

mort = [
    (D(2026, 7, 4), "P006", "BROILER-JUL-001", 4, "Unknown"), (D(2026, 7, 15), "P006", "BROILER-JUL-001", 3, "Heat stress"),
    (D(2026, 8, 2), "P006", "BROILER-JUL-001", 3, "Smothering / Crushing"),
    (D(2026, 8, 20), "P005", "LAYER-JUN-001", 2, "Unknown"),
    (D(2026, 9, 3), "P001", "BROILER-SEP-001", 2, "Unknown"), (D(2026, 9, 4), "P002", "BROILER-SEP-002", 3, "Cold / Chilling"),
    (D(2026, 9, 9), "P003", "BROILER-SEP-003", 1, "Injury"), (D(2026, 9, 14), "P002", "BROILER-SEP-002", 2, "Suspected disease"),
    (D(2026, 9, 20), "P004", "BROILER-SEP-004", 2, "Unknown"), (D(2026, 9, 27), "P001", "BROILER-SEP-001", 1, "Heat stress"),
]
mort_rows = [dict(zip(["Date", "Pen ID", "Batch ID", "Number of Deaths", "Suspected/Recorded Cause"], m), Notes=S)
             for m in mort]

prod = [
    (D(2026, 8, 15), "P006", "BROILER-JUL-001", "Weighing", None, "kg", 2.35, 20),
    (D(2026, 9, 10), "P005", "LAYER-JUN-001", "Egg Collection", 11, "crates", None, None),
    (D(2026, 9, 14), "P001", "BROILER-SEP-001", "Weighing", None, "kg", 0.45, 20),
    (D(2026, 9, 14), "P002", "BROILER-SEP-002", "Weighing", None, "kg", 0.43, 20),
    (D(2026, 9, 20), "P005", "LAYER-JUN-001", "Egg Collection", 11, "crates", None, None),
    (D(2026, 9, 21), "P001", "BROILER-SEP-001", "Weighing", None, "kg", 0.86, 20),
    (D(2026, 9, 21), "P004", "BROILER-SEP-004", "Weighing", None, "kg", 0.80, 25),
    (D(2026, 9, 27), "P005", "LAYER-JUN-001", "Egg Collection", 10, "crates", None, None),
]
prod_rows = [dict(zip(["Date", "Pen ID", "Batch ID", "Measurement Type", "Quantity", "Unit", "Average Weight kg",
                       "Number Sampled"], p), Notes=S) for p in prod]

sales = [
    (D(2026, 8, 18), "P006", "BROILER-JUL-001", "Live Birds", 290, 5500, "Market Trader (sample)", "Transfer", "SL-001"),
    (D(2026, 9, 12), "P005", "LAYER-JUN-001", "Eggs", 20, 4800, "Customer A (sample)", "Cash", "SL-002"),
    (D(2026, 9, 20), "GENERAL", None, "Manure", 10, 1000, "Walk-in Customer", "Cash", "SL-003"),
    (D(2026, 9, 25), "P005", "LAYER-JUN-001", "Eggs", 18, 4900, "Customer B (sample)", "Credit", "SL-004"),
]
sales_rows = [dict(zip(["Date", "Pen ID", "Batch ID", "Product", "Quantity Sold", "Unit Price", "Customer",
                        "Payment Method", "Reference"], s), Notes=S) for s in sales]

exps = [
    (D(2026, 6, 1), "Day-old Chicks / Stock Purchase", "400 point-of-lay pullets", "P005", "LAYER-JUN-001", 1400000, "Transfer"),
    (D(2026, 7, 1), "Day-old Chicks / Stock Purchase", "300 day-old chicks", "P006", "BROILER-JUL-001", 210000, "Transfer"),
    (D(2026, 7, 1), "Litter / Bedding", "Wood shavings", "P006", "BROILER-JUL-001", 12000, "Cash"),
    (D(2026, 7, 31), "Labour", "July wages - all pens", "GENERAL", None, 150000, "Transfer"),
    (D(2026, 8, 18), "Transport", "Delivery of birds to market", "P006", "BROILER-JUL-001", 15000, "Cash"),
    (D(2026, 8, 31), "Labour", "August wages - all pens", "GENERAL", None, 150000, "Transfer"),
    (D(2026, 9, 1), "Day-old Chicks / Stock Purchase", "500 day-old chicks", "P001", "BROILER-SEP-001", 425000, "Transfer"),
    (D(2026, 9, 1), "Day-old Chicks / Stock Purchase", "600 day-old chicks", "P002", "BROILER-SEP-002", 510000, "Transfer"),
    (D(2026, 9, 2), "Day-old Chicks / Stock Purchase", "450 day-old chicks", "P003", "BROILER-SEP-003", 382500, "Transfer"),
    (D(2026, 9, 2), "Day-old Chicks / Stock Purchase", "550 day-old chicks", "P004", "BROILER-SEP-004", 467500, "Transfer"),
    (D(2026, 9, 1), "Litter / Bedding", "Wood shavings", "P001", "BROILER-SEP-001", 12000, "Cash"),
    (D(2026, 9, 11), "Repairs", "Roof repair - P002 (pen-level, no batch)", "P002", None, 18000, "Cash"),
    (D(2026, 9, 26), "Electricity", "September electricity bill", "GENERAL", None, 35000, "Transfer"),
]
exp_rows = [dict(zip(["Date", "Category", "Description", "Pen ID / General", "Batch ID (if applicable)", "Amount",
                      "Payment Method"], e), Notes=S) for e in exps]

moves = [
    (D(2026, 9, 5), "P002", "BROILER-SEP-002", "Transfer Out", None, 10, "BROILER-SEP-001"),
    (D(2026, 9, 5), "P001", "BROILER-SEP-001", "Transfer In", 10, None, "BROILER-SEP-002"),
]
move_rows = [dict(zip(["Date", "Pen ID", "Batch ID", "Movement Type", "Quantity In", "Quantity Out",
                       "Related Batch ID"], m), Notes=S + ". Overcrowding relief.") for m in moves]

# cashbook sample built from the non-credit sample transactions so it reconciles
cash = [(D(2026, 5, 31), "Opening Balance", "", "Opening cash & bank balance", "GENERAL", 5000000, None, "Transfer", "OB")]
for f in feed_pur:
    if f[6] != "Credit":
        cash.append((f[0], "Feed Purchase", "Feed", f"{f[3]} bags {f[1]}", "GENERAL", None, f[3] * f[5], f[6], f[7]))
for d_ in drug_pur:
    if d_[7] != "Credit":
        cash.append((d_[0], "Drug Purchase", "Drugs", f"{d_[1]}", "GENERAL", None, d_[3] * d_[4], d_[7], d_[6]))
for e in exps:
    if e[6] != "Credit":
        cash.append((e[0], "Expense Payment", e[1], e[2], e[3], None, e[5], e[6], ""))
for s in sales:
    if s[7] != "Credit":
        cash.append((s[0], "Sale Receipt", "Sales", f"{s[4]} {s[3]}", s[1], s[4] * s[5], None, s[7], s[8]))
cash.append((D(2026, 9, 27), "Supplier Credit Payment", "Feed", "Part payment for INV-1102 (Grower on credit)",
             "GENERAL", None, 200000, "Transfer", "INV-1102"))
cash.sort(key=lambda x: x[0])
cash_rows = [dict(zip(["Date", "Type", "Category", "Description", "Pen ID / General", "Money In", "Money Out",
                       "Payment Method", "Reference"], c)) for c in cash]

# --------------------------------------------------------------------------------------
# PEN_REGISTER
# --------------------------------------------------------------------------------------
T = "tPens"
active_cnt = f"COUNTIFS({tr('tBatches', 'Pen ID')},{me(T, 'Pen ID')},{tr('tBatches', 'Status')},\"Active\")"
pen_cols = [
    Col("Pen ID", 10, TXT, note="Permanent, unique ID e.g. P001. Never reuse or rename once records exist."),
    Col("Pen Name", 18),
    Col("Farm/Location", 14),
    Col("Animal Type", 12, dv=dv_list("L_Animals")),
    Col("Current Batch ID", 20, formula=f'IF({active_cnt}=0,"",IF({active_cnt}=1,'
        f'INDEX({tr("tBatches", "Batch ID")},MATCH({me(T, "Pen ID")},{tr("tBatches", "Active Pen Key")},0)),'
        f'INDEX({tr("tBatches", "Batch ID")},MATCH({me(T, "Pen ID")},{tr("tBatches", "Active Pen Key")},0))'
        f'&" +"&({active_cnt}-1)&" more"))'),
    Col("Start Date", 12, DMY, formula=f'IF({active_cnt}=0,"",_xlfn.MINIFS({tr("tBatches", "Batch Start Date")},'
        f'{tr("tBatches", "Pen ID")},{me(T, "Pen ID")},{tr("tBatches", "Status")},"Active"))'),
    Col("Initial Stock", 11, INT, formula=f'SUMIFS({tr("tBatches", "Initial Quantity")},{tr("tBatches", "Pen ID")},'
        f'{me(T, "Pen ID")},{tr("tBatches", "Status")},"Active")'),
    Col("Current Stock", 11, INT, formula=f'SUMIFS({tr("tBatches", "Current Quantity")},{tr("tBatches", "Pen ID")},'
        f'{me(T, "Pen ID")},{tr("tBatches", "Status")},"Active")'),
    Col("Status", 10, formula=f'IF({me(T, "Pen ID")}="","",IF({me(T, "Pen Closed")}="Yes","Closed",'
        f'IF({active_cnt}>0,"Active","Empty")))'),
    Col("Pen Closed", 10, dv=dv_list('"No,Yes"'), note="Set to Yes only if the pen is permanently out of use."),
    Col("Notes", 40),
    Col("Check", 30, formula=blank_guard(T, ["Pen ID"], nested_if([
        (f"COUNTIF({tr(T, 'Pen ID')},{me(T, 'Pen ID')})>1", '"Duplicate Pen ID"'),
        (f'{me(T, "Pen ID")}="GENERAL"', '"GENERAL is reserved - use another ID"'),
        (f'AND({me(T, "Pen Closed")}="Yes",{active_cnt}>0)', '"Pen marked closed but has an active batch"'),
    ], '"OK"'))),
]
ws = WS["PEN_REGISTER"]
sheet_title(ws, "PEN REGISTER", "Physical pens (permanent). Current batch, stock and status are calculated from "
            "BATCH_REGISTER. A pen keeps its ID for life; batches come and go.")
build_table(ws, T, pen_cols, pens)
check_cf(ws, col_letter(pen_cols, "Check"))
L = col_letter(pen_cols, "Status")
for val, fill in (("Active", FILL_GREEN), ("Empty", FILL_AMBER), ("Closed", PatternFill("solid", fgColor="D9D9D9"))):
    ws.conditional_formatting.add(f"{L}{FIRST}:{L}{MAX_ROW}",
                                  FormulaRule(formula=[f'{L}{FIRST}="{val}"'], fill=fill))

# --------------------------------------------------------------------------------------
# BATCH_REGISTER
# --------------------------------------------------------------------------------------
T = "tBatches"
b_id = me(T, "Batch ID")
batch_cols = [
    Col("Batch ID", 18, TXT, note="Unique ID e.g. BROILER-SEP-001. Never reuse."),
    Col("Pen ID", 9, dv=dv_list("L_Pens")),
    Col("Animal Type", 11, dv=dv_list("L_Animals")),
    Col("Batch Start Date", 12, DMY, dv=DV_DATE),
    Col("Initial Quantity", 10, INT, dv=DV_INT),
    Col("Source/Supplier", 18, dv=dv_list("L_Suppliers", False)),
    Col("Breed/Type", 16),
    Col("Current Quantity", 10, INT, formula=f"{me(T, 'Total Placed')}-{me(T, 'Deaths')}-{me(T, 'Heads Sold')}"
        f"-{me(T, 'Transfers Out')}"),
    Col("Closing Date", 12, DMY, dv=DV_DATE),
    Col("Status", 9, dv=dv_list('"Active,Closed"')),
    Col("Notes", 34),
    Col("Transfers In / Added", 11, INT, formula=f"SUMIFS({tr('tMoves', 'Quantity In')},{tr('tMoves', 'Batch ID')},{b_id})"),
    Col("Deaths", 9, INT, formula=f"SUMIFS({tr('tMort', 'Number of Deaths')},{tr('tMort', 'Batch ID')},{b_id})"),
    Col("Heads Sold", 9, INT, formula=f"SUMIFS({tr('tSales', 'Heads Removed')},{tr('tSales', 'Batch ID')},{b_id})"),
    Col("Transfers Out", 10, INT, formula=f"SUMIFS({tr('tMoves', 'Quantity Out')},{tr('tMoves', 'Batch ID')},{b_id})"),
    Col("Total Placed", 10, INT, formula=f"N({me(T, 'Initial Quantity')})+{me(T, 'Transfers In / Added')}"),
    Col("Age (days)", 8, INT, formula=f'IF(NOT(ISNUMBER({me(T, "Batch Start Date")})),"",IF(N({me(T, "Closing Date")})>0,'
        f'{me(T, "Closing Date")},TODAY())-{me(T, "Batch Start Date")}+1)'),
    Col("Active Pen Key", 9, formula=f'IF({me(T, "Status")}="Active",{me(T, "Pen ID")},"")'),
    Col("Check", 34, formula=blank_guard(T, ["Batch ID", "Pen ID"], nested_if([
        (f'{b_id}=""', '"Missing Batch ID"'),
        (f"COUNTIF({tr(T, 'Batch ID')},{b_id})>1", '"Duplicate Batch ID"'),
        (f"COUNTIF({tr('tPens', 'Pen ID')},{me(T, 'Pen ID')})=0", '"Pen is not in PEN_REGISTER"'),
        (f"NOT(ISNUMBER({me(T, 'Batch Start Date')}))", '"Missing start date"'),
        (f"N({me(T, 'Initial Quantity')})<=0", '"Initial quantity must be > 0"'),
        (f'{me(T, "Status")}=""', '"Missing status"'),
        (f"{me(T, 'Current Quantity')}<0", '"Negative stock - more deaths/sales than animals"'),
        (f'AND({me(T, "Status")}="Closed",N({me(T, "Closing Date")})=0)', '"Closed batch needs a closing date"'),
        (f'AND({me(T, "Status")}="Closed",{me(T, "Current Quantity")}>0)',
         f'"Closed with "&{me(T, "Current Quantity")}&" animals unaccounted for"'),
        (f"AND(N({me(T, 'Closing Date')})>0,N({me(T, 'Closing Date')})<{me(T, 'Batch Start Date')})",
         '"Closing date is before start date"'),
    ], '"OK"'))),
]
ws = WS["BATCH_REGISTER"]
sheet_title(ws, "BATCH REGISTER", "One row per production cycle. A batch lives in one pen. When finished set Status = "
            "Closed and enter the Closing Date - its history is kept for ever and the pen becomes free for a new batch.")
build_table(ws, T, batch_cols, batch_rows)
check_cf(ws, col_letter(batch_cols, "Check"))
L = col_letter(batch_cols, "Current Quantity")
ws.conditional_formatting.add(f"{L}{FIRST}:{L}{MAX_ROW}", CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED))

# --------------------------------------------------------------------------------------
# FEED_PURCHASES
# --------------------------------------------------------------------------------------
T = "tFeedPur"
fp_cols = [
    Col("Date", 11, DMY, dv=DV_DATE),
    Col("Purchase ID", 11, formula=f'"FP-"&TEXT(ROW()-ROW({T}[[#Headers],[Date]]),"00000")'),
    Col("Feed Type", 13, dv=dv_list("L_FeedTypes")),
    Col("Supplier", 20, dv=dv_list("L_Suppliers", False)),
    Col("Quantity in Bags", 9, DEC1, dv=DV_NUM),
    Col("Bag Size in kg", 8, DEC1, dv=DV_NUM),
    Col("Total Quantity in kg", 10, DEC1, formula=f"N({me(T, 'Quantity in Bags')})*N({me(T, 'Bag Size in kg')})"),
    Col("Price per Bag", 11, NGN, dv=DV_NUM),
    Col("Total Cost", 13, NGN, formula=f"N({me(T, 'Quantity in Bags')})*N({me(T, 'Price per Bag')})"),
    Col("Cost per kg", 10, NGN2, formula=f"IFERROR({me(T, 'Total Cost')}/{me(T, 'Total Quantity in kg')},0)"),
    Col("Payment Method", 10, dv=dv_list("L_Pay")),
    Col("Reference", 11),
    Col("Notes", 34),
] + moving_avg_cols(T, "Feed Type", "Total Quantity in kg", "Total Cost", "tFeedIss", "Feed Type",
                    "Quantity Issued in kg", "Purchase No. of Type", "Stock Before Purchase kg",
                    "Avg Cost per kg After Purchase", DEC1, NGN2) + [
    Col("Check", 26, formula=blank_guard(T, ["Date", "Feed Type", "Quantity in Bags"], nested_if(
        date_checks(T) + [order_check(T),
            (f'{me(T, "Feed Type")}=""', '"Missing feed type"'),
            (f"COUNTIF({tr('tFeedTypes', 'Feed Type')},{me(T, 'Feed Type')})=0", '"Feed type not in LISTS"'),
            (f"N({me(T, 'Quantity in Bags')})<=0", '"Bags must be greater than 0"'),
            (f"N({me(T, 'Bag Size in kg')})<=0", '"Missing bag size"'),
            (f'{me(T, "Price per Bag")}=""', '"Missing price"'),
            (f'{me(T, "Payment Method")}=""', '"Missing payment method"'),
        ], '"OK"'))),
]
ws = WS["FEED_PURCHASES"]
sheet_title(ws, "FEED PURCHASES", "Record each purchase ONCE. Feed goes into the store (INVENTORY) - it only becomes "
            "a pen/batch cost when it is issued on FEED_ISSUES. Inventory unit = kg.")
build_table(ws, T, fp_cols, feed_pur_rows)
check_cf(ws, col_letter(fp_cols, "Check"))

# --------------------------------------------------------------------------------------
# FEED_ISSUES
# --------------------------------------------------------------------------------------
T = "tFeedIss"
ft, dt = me(T, "Feed Type"), me(T, "Date")
wac_feed = latest_avg("tFeedPur", "Feed Type", ft, dt, "Avg Cost per kg After Purchase")
fi_cols = [
    Col("Date", 11, DMY, dv=DV_DATE),
    Col("Issue ID", 11, formula=f'"FI-"&TEXT(ROW()-ROW({T}[[#Headers],[Date]]),"00000")'),
    Col("Feed Type", 13, dv=dv_list("L_FeedTypes")),
    Col("Pen ID", 8, dv=dv_list("L_Pens")),
    Col("Batch ID", 18, dv=dv_list("L_Batches")),
    Col("Quantity Issued in kg", 10, DEC1, dv=DV_NUM),
    Col("Cost per kg", 10, NGN2, formula=wac_feed,
        note="Moving weighted-average cost of this feed type in store on the issue date (from FEED_PURCHASES)."),
    Col("Total Cost", 12, NGN, formula=f"N({me(T, 'Quantity Issued in kg')})*{me(T, 'Cost per kg')}"),
    Col("Issued By", 13, dv=dv_list("L_Staff", False)),
    Col("Notes", 34),
    Col("Animal Type", 10, formula=animal_lookup(T)),
    Col("Store Balance on Date kg", 11, DEC1,
        formula=f"SUMIFS({tr('tFeedPur', 'Total Quantity in kg')},{tr('tFeedPur', 'Feed Type')},{ft},"
                f"{tr('tFeedPur', 'Date')},\"<=\"&{dt})-SUMIFS({tr(T, 'Quantity Issued in kg')},{tr(T, 'Feed Type')},"
                f"{ft},{tr(T, 'Date')},\"<=\"&{dt})",
        note="Feed of this type left in store at the end of this date after all issues on or before it."),
    Col("Check", 30, formula=blank_guard(T, ["Date", "Feed Type", "Pen ID", "Batch ID", "Quantity Issued in kg"],
        nested_if(date_checks(T) + [
            (f'{ft}=""', '"Missing feed type"'),
            (f"COUNTIF({tr('tFeedTypes', 'Feed Type')},{ft})=0", '"Feed type not in LISTS"'),
        ] + pen_batch_checks(T) + [
            (f"N({me(T, 'Quantity Issued in kg')})<=0", '"Quantity must be greater than 0"'),
            (f"{me(T, 'Cost per kg')}=0", '"No purchase of this feed on/before this date"'),
            (f"{me(T, 'Store Balance on Date kg')}<-0.0001",
             f'"EXCEEDS STOCK by "&TEXT(-{me(T, "Store Balance on Date kg")},"#,##0.0")&" kg"'),
        ], '"OK"'))),
]
ws = WS["FEED_ISSUES"]
sheet_title(ws, "FEED ISSUES", "Each time feed leaves the store for a pen, record it here. This reduces inventory "
            "and charges the cost (weighted-average) to the pen & batch. Record only what happened.")
build_table(ws, T, fi_cols, feed_iss_rows)
check_cf(ws, col_letter(fi_cols, "Check"))

# --------------------------------------------------------------------------------------
# DRUG_PURCHASES
# --------------------------------------------------------------------------------------
T = "tDrugPur"
dn = me(T, "Drug Name")
lot_rem = (f'IF({me(T, "Batch Number")}="","",N({me(T, "Quantity Purchased")})-SUMIFS({tr("tDrugUse", "Quantity Used")},'
           f'{tr("tDrugUse", "Drug Name")},{dn},{tr("tDrugUse", "Drug Batch Number")},{me(T, "Batch Number")}))')
dp_cols = [
    Col("Date", 11, DMY, dv=DV_DATE),
    Col("Purchase ID", 11, formula=f'"DP-"&TEXT(ROW()-ROW({T}[[#Headers],[Date]]),"00000")'),
    Col("Drug Name", 20, dv=dv_list("L_Drugs")),
    Col("Form/Concentration", 18, formula=f'IFERROR(INDEX({tr("tDrugs", "Form/Concentration")},'
        f'MATCH({dn},{tr("tDrugs", "Drug Name")},0))&"","")'),
    Col("Supplier", 18, dv=dv_list("L_Suppliers", False)),
    Col("Quantity Purchased", 10, DEC1, dv=DV_NUM, note="Enter in the drug's Stock Unit (see Unit column / LISTS)."),
    Col("Unit", 8, formula=f'IFERROR(INDEX({tr("tDrugs", "Stock Unit")},MATCH({dn},{tr("tDrugs", "Drug Name")},0))&"","")'),
    Col("Unit Cost", 10, NGN2, dv=DV_NUM),
    Col("Total Cost", 12, NGN, formula=f"N({me(T, 'Quantity Purchased')})*N({me(T, 'Unit Cost')})"),
    Col("Expiry Date", 11, DMY, dv=DV_DATE),
    Col("Batch Number", 11, TXT, note="Manufacturer lot/batch number. Quote it on DRUG_USAGE to track the lot."),
    Col("Payment Method", 10, dv=dv_list("L_Pay")),
    Col("Notes", 30),
    Col("Remaining in Lot", 10, DEC1, formula=lot_rem),
] + moving_avg_cols(T, "Drug Name", "Quantity Purchased", "Total Cost", "tDrugUse", "Drug Name", "Quantity Used",
                    "Purchase No. of Drug", "Stock Before Purchase", "Avg Cost per Unit After Purchase",
                    DEC1, NGN2) + [
    Col("Expiry Status", 16, formula=f'IF(N({me(T, "Expiry Date")})=0,"",IF(AND({me(T, "Remaining in Lot")}<>"",'
        f'N({me(T, "Remaining in Lot")})<=0),"Used up",IF({me(T, "Expiry Date")}<TODAY(),"EXPIRED",'
        f'IF({me(T, "Expiry Date")}-TODAY()<=ExpiryWarnDays,"Expires in "&({me(T, "Expiry Date")}-TODAY())&" days","OK"))))'),
    Col("Check", 26, formula=blank_guard(T, ["Date", "Drug Name", "Quantity Purchased"], nested_if(date_checks(T) + [
        order_check(T),
        (f'{dn}=""', '"Missing drug name"'),
        (f"COUNTIF({tr('tDrugs', 'Drug Name')},{dn})=0", '"Drug not in LISTS"'),
        (f"N({me(T, 'Quantity Purchased')})<=0", '"Quantity must be greater than 0"'),
        (f'{me(T, "Unit Cost")}=""', '"Missing unit cost"'),
        (f"N({me(T, 'Expiry Date')})=0", '"Missing expiry date"'),
        (f'{me(T, "Payment Method")}=""', '"Missing payment method"'),
    ], '"OK"'))),
]
ws = WS["DRUG_PURCHASES"]
sheet_title(ws, "DRUG PURCHASES", "Record each drug/vaccine purchase once. Unused drugs remain inventory. Quantities "
            "are in the drug's Stock Unit set on LISTS (ml, g, doses, tablets ...).")
build_table(ws, T, dp_cols, drug_pur_rows)
check_cf(ws, col_letter(dp_cols, "Check"))
L = col_letter(dp_cols, "Expiry Status")
ws.conditional_formatting.add(f"{L}{FIRST}:{L}{MAX_ROW}", FormulaRule(formula=[f'{L}{FIRST}="EXPIRED"'],
                              fill=FILL_RED, font=Font(name=FONT, bold=True, color=RED_DARK)))
ws.conditional_formatting.add(f"{L}{FIRST}:{L}{MAX_ROW}", FormulaRule(formula=[f'LEFT({L}{FIRST},7)="Expires"'],
                              fill=FILL_AMBER))

# --------------------------------------------------------------------------------------
# DRUG_USAGE
# --------------------------------------------------------------------------------------
T = "tDrugUse"
dn, dt, lot = me(T, "Drug Name"), me(T, "Date"), me(T, "Drug Batch Number")
wac_drug = latest_avg("tDrugPur", "Drug Name", dn, dt, "Avg Cost per Unit After Purchase")
lot_exp = f"_xlfn.MAXIFS({tr('tDrugPur', 'Expiry Date')},{tr('tDrugPur', 'Drug Name')},{dn},{tr('tDrugPur', 'Batch Number')},{lot})"
du_cols = [
    Col("Date", 11, DMY, dv=DV_DATE),
    Col("Usage ID", 11, formula=f'"DU-"&TEXT(ROW()-ROW({T}[[#Headers],[Date]]),"00000")'),
    Col("Drug Name", 20, dv=dv_list("L_Drugs")),
    Col("Pen ID", 8, dv=dv_list("L_Pens")),
    Col("Batch ID", 18, dv=dv_list("L_Batches")),
    Col("Reason", 14, dv=dv_list("L_Reasons", False)),
    Col("Dosage", 20, note="Free text, e.g. 10 ml per 20 litres"),
    Col("Quantity Used", 10, DEC1, dv=DV_NUM, note="Total quantity taken from stock, in the drug's Stock Unit."),
    Col("Unit", 8, formula=f'IFERROR(INDEX({tr("tDrugs", "Stock Unit")},MATCH({dn},{tr("tDrugs", "Drug Name")},0))&"","")'),
    Col("Cost per Unit", 10, NGN2, formula=wac_drug,
        note="Moving weighted-average cost of this drug in store on the usage date (from DRUG_PURCHASES)."),
    Col("Total Cost", 11, NGN, formula=f"N({me(T, 'Quantity Used')})*{me(T, 'Cost per Unit')}"),
    Col("Drug Batch Number", 11, TXT, dv=None),
    Col("Administered By", 13, dv=dv_list("L_Staff", False)),
    Col("Notes", 30),
    Col("Animal Type", 10, formula=animal_lookup(T)),
    Col("Store Balance on Date", 10, DEC1,
        formula=f"SUMIFS({tr('tDrugPur', 'Quantity Purchased')},{tr('tDrugPur', 'Drug Name')},{dn},"
                f"{tr('tDrugPur', 'Date')},\"<=\"&{dt})-SUMIFS({tr(T, 'Quantity Used')},{tr(T, 'Drug Name')},{dn},"
                f"{tr(T, 'Date')},\"<=\"&{dt})"),
    Col("Check", 30, formula=blank_guard(T, ["Date", "Drug Name", "Pen ID", "Batch ID", "Quantity Used"],
        nested_if(date_checks(T) + [
            (f'{dn}=""', '"Missing drug name"'),
            (f"COUNTIF({tr('tDrugs', 'Drug Name')},{dn})=0", '"Drug not in LISTS"'),
        ] + pen_batch_checks(T) + [
            (f"N({me(T, 'Quantity Used')})<=0", '"Quantity must be greater than 0"'),
            (f"{me(T, 'Cost per Unit')}=0", '"No purchase of this drug on/before this date"'),
            (f"{me(T, 'Store Balance on Date')}<-0.0001",
             f'"EXCEEDS STOCK by "&TEXT(-{me(T, "Store Balance on Date")},"#,##0.0")&" "&{me(T, "Unit")}'),
            (f"AND({lot}<>\"\",COUNTIFS({tr('tDrugPur', 'Drug Name')},{dn},{tr('tDrugPur', 'Batch Number')},{lot})=0)",
             '"Drug batch number not found in purchases"'),
            (f"AND({lot}<>\"\",N({lot_exp})>0,{dt}>{lot_exp})", f'"Lot was expired on "&TEXT({lot_exp},"dd/mm/yyyy")'),
        ], '"OK"'))),
]
ws = WS["DRUG_USAGE"]
sheet_title(ws, "DRUG USAGE", "One row per administration per pen/batch. The same drug used in several pens at "
            "different dosages = separate rows. Cost is charged to the batch at weighted-average cost.")
build_table(ws, T, du_cols, drug_use_rows)
check_cf(ws, col_letter(du_cols, "Check"))

# --------------------------------------------------------------------------------------
# MORTALITY_LOG
# --------------------------------------------------------------------------------------
T = "tMort"
mo_cols = [
    Col("Date", 11, DMY, dv=DV_DATE),
    Col("Mortality ID", 11, formula=f'"MO-"&TEXT(ROW()-ROW({T}[[#Headers],[Date]]),"00000")'),
    Col("Pen ID", 8, dv=dv_list("L_Pens")),
    Col("Batch ID", 18, dv=dv_list("L_Batches")),
    Col("Number of Deaths", 9, INT, dv=DV_INT),
    Col("Suspected/Recorded Cause", 20, dv=dv_list("L_Causes", False),
        note="Record only what was observed/diagnosed. Use Unknown when unsure."),
    Col("Age/Day of Batch", 9, INT, formula=f'IFERROR({me(T, "Date")}-{bl("Batch Start Date", me(T, "Batch ID"))}+1,"")'),
    Col("Notes", 34),
    Col("Animal Type", 10, formula=animal_lookup(T)),
    Col("Check", 30, formula=blank_guard(T, ["Date", "Pen ID", "Batch ID", "Number of Deaths"],
        nested_if(date_checks(T) + pen_batch_checks(T) + [
            (f"N({me(T, 'Number of Deaths')})<=0", '"Number of deaths must be > 0"'),
            (f"{bl('Current Quantity', me(T, 'Batch ID'))}<0", '"Batch stock is negative - check counts"'),
        ], '"OK"'))),
]
ws = WS["MORTALITY_LOG"]
sheet_title(ws, "MORTALITY LOG", "Only record days on which animals died. No entry = no deaths. Current stock "
            "updates automatically.")
build_table(ws, T, mo_cols, mort_rows)
check_cf(ws, col_letter(mo_cols, "Check"))

# --------------------------------------------------------------------------------------
# PRODUCTION_LOG
# --------------------------------------------------------------------------------------
T = "tProd"
pr_cols = [
    Col("Date", 11, DMY, dv=DV_DATE),
    Col("Record ID", 11, formula=f'"PR-"&TEXT(ROW()-ROW({T}[[#Headers],[Date]]),"00000")'),
    Col("Pen ID", 8, dv=dv_list("L_Pens")),
    Col("Batch ID", 18, dv=dv_list("L_Batches")),
    Col("Measurement Type", 14, dv=dv_list("L_Meas", False)),
    Col("Quantity", 9, DEC1, dv=DV_NUM, note="Output quantity, e.g. crates of eggs collected. Leave blank for weighing."),
    Col("Unit", 8, dv=dv_list("L_Units", False)),
    Col("Average Weight kg", 10, WT, dv=DV_NUM, note="Average live weight per animal (kg) for Weighing records."),
    Col("Number Sampled", 9, INT, dv=DV_INT),
    Col("Notes", 30),
    Col("Animal Type", 10, formula=animal_lookup(T)),
    Col("Age/Day of Batch", 9, INT, formula=f'IFERROR({me(T, "Date")}-{bl("Batch Start Date", me(T, "Batch ID"))}+1,"")'),
    Col("Check", 30, formula=blank_guard(T, ["Date", "Pen ID", "Batch ID", "Measurement Type"],
        nested_if(date_checks(T) + pen_batch_checks(T) + [
            (f'{me(T, "Measurement Type")}=""', '"Missing measurement type"'),
            (f'AND({me(T, "Measurement Type")}="Weighing",N({me(T, "Average Weight kg")})<=0)',
             '"Weighing needs an average weight"'),
            (f'AND({me(T, "Measurement Type")}<>"Weighing",N({me(T, "Quantity")})<=0)', '"Output needs a quantity"'),
        ], '"OK"'))),
]
ws = WS["PRODUCTION_LOG"]
sheet_title(ws, "PRODUCTION LOG", "Weighings and output (eggs etc.) - only when they happen. No daily row needed.")
build_table(ws, T, pr_cols, prod_rows)
check_cf(ws, col_letter(pr_cols, "Check"))

# --------------------------------------------------------------------------------------
# SALES
# --------------------------------------------------------------------------------------
T = "tSales"
prd = me(T, "Product")
reduces = f'IFERROR(INDEX({tr("tProducts", "Reduces Head Count")},MATCH({prd},{tr("tProducts", "Product")},0)),"No")'
p, b = me(T, "Pen ID"), me(T, "Batch ID")
sa_cols = [
    Col("Date", 11, DMY, dv=DV_DATE),
    Col("Sale ID", 11, formula=f'"SA-"&TEXT(ROW()-ROW({T}[[#Headers],[Date]]),"00000")'),
    Col("Pen ID", 9, dv=dv_list("L_PensGen"), note="Pen the product came from, or GENERAL for farm-wide products."),
    Col("Batch ID", 18, dv=dv_list("L_Batches")),
    Col("Product", 13, dv=dv_list("L_Products")),
    Col("Quantity Sold", 9, DEC1, dv=DV_NUM),
    Col("Unit Price", 11, NGN, dv=DV_NUM),
    Col("Total Revenue", 13, NGN, formula=f"N({me(T, 'Quantity Sold')})*N({me(T, 'Unit Price')})"),
    Col("Customer", 18, dv=dv_list("L_Customers", False)),
    Col("Payment Method", 10, dv=dv_list("L_Pay")),
    Col("Reference", 10),
    Col("Notes", 26),
    Col("Sale Unit", 8, formula=f'IFERROR(INDEX({tr("tProducts", "Sale Unit")},MATCH({prd},{tr("tProducts", "Product")},0))&"","")'),
    Col("Heads Removed", 9, INT, formula=f'IF({reduces}="Yes",N({me(T, "Quantity Sold")}),0)',
        note="Animals leaving the batch because of this sale (only for products marked 'Reduces Head Count')."),
    Col("Animal Type", 10, formula=animal_lookup(T)),
    Col("Check", 30, formula=blank_guard(T, ["Date", "Pen ID", "Batch ID", "Product", "Quantity Sold"],
        nested_if(date_checks(T) + [
            (f'{prd}=""', '"Missing product"'),
            (f"COUNTIF({tr('tProducts', 'Product')},{prd})=0", '"Product not in LISTS"'),
            (f"N({me(T, 'Quantity Sold')})<=0", '"Quantity must be greater than 0"'),
            (f'{me(T, "Unit Price")}=""', '"Missing unit price"'),
            (f'{p}=""', '"Missing Pen ID (or GENERAL)"'),
            (f'AND({p}="GENERAL",{b}<>"")', '"GENERAL sales cannot carry a batch"'),
            (f'AND({p}="GENERAL",{reduces}="Yes")', '"Animal sales must name the pen & batch"'),
            (f'AND({p}<>"GENERAL",COUNTIF({tr("tPens", "Pen ID")},{p})=0)', f'"Pen "&{p}&" is not in PEN_REGISTER"'),
            (f'AND({b}="",{reduces}="Yes")', '"Missing Batch ID"'),
            (f'AND({b}<>"",COUNTIF({tr("tBatches", "Batch ID")},{b})=0)', '"Batch is not in BATCH_REGISTER"'),
            (f'AND({b}<>"",COUNTIFS({tr("tBatches", "Batch ID")},{b},{tr("tBatches", "Pen ID")},{p})=0)',
             f'"Batch "&{b}&" belongs to pen "&{bl("Pen ID", b)}'),
            (f'IFERROR(AND({b}<>"",{me(T, "Date")}<{bl("Batch Start Date", b)}),FALSE)', '"Date is before the batch start date"'),
            (f'IFERROR(AND({b}<>"",{bl("Current Quantity", b)}<0),FALSE)', '"Sold more animals than the batch holds"'),
            (f'{me(T, "Payment Method")}=""', '"Missing payment method"'),
        ], '"OK"'))),
]
ws = WS["SALES"]
sheet_title(ws, "SALES", "Every sale, linked to the pen & batch it came from (GENERAL for farm-wide items such as "
            "manure). Animal sales reduce the batch's current stock automatically.")
build_table(ws, T, sa_cols, sales_rows)
check_cf(ws, col_letter(sa_cols, "Check"))

# --------------------------------------------------------------------------------------
# OTHER_EXPENSES
# --------------------------------------------------------------------------------------
T = "tExp"
p, b = me(T, "Pen ID / General"), me(T, "Batch ID (if applicable)")
ex_cols = [
    Col("Date", 11, DMY, dv=DV_DATE),
    Col("Expense ID", 11, formula=f'"EX-"&TEXT(ROW()-ROW({T}[[#Headers],[Date]]),"00000")'),
    Col("Category", 22, dv=dv_list("L_ExpCat")),
    Col("Description", 30),
    Col("Pen ID / General", 10, dv=dv_list("L_PensGen"), note="Pen ID for a pen-specific cost, GENERAL for farm-wide."),
    Col("Batch ID (if applicable)", 18, dv=dv_list("L_Batches"),
        note="Fill in to include this cost in the batch's profitability (e.g. chicks, litter)."),
    Col("Amount", 12, NGN, dv=DV_NUM),
    Col("Payment Method", 10, dv=dv_list("L_Pay")),
    Col("Reference", 10),
    Col("Notes", 26),
    Col("Cost Level", 16, formula=f'IF({p}="","",IF({p}="GENERAL","General overhead",IF({b}="","Pen-level","Batch direct")))',
        note="Batch direct = in BATCH_PROFITABILITY. Pen-level = in pen totals only. General = farm totals only."),
    Col("Animal Type", 10, formula=f'IF({b}<>"",IFERROR({bl("Animal Type", b)},""),IFERROR(INDEX({tr("tPens", "Animal Type")},'
        f'MATCH({p},{tr("tPens", "Pen ID")},0))&"",""))'),
    Col("Check", 30, formula=blank_guard(T, ["Date", "Category", "Pen ID / General", "Amount"],
        nested_if(date_checks(T) + [
            (f'{me(T, "Category")}=""', '"Missing category"'),
            (f"N({me(T, 'Amount')})<=0", '"Amount must be greater than 0"'),
            (f'{p}=""', '"Enter a Pen ID or GENERAL"'),
            (f'AND({p}<>"GENERAL",COUNTIF({tr("tPens", "Pen ID")},{p})=0)', f'"Pen "&{p}&" is not in PEN_REGISTER"'),
            (f'AND({p}="GENERAL",{b}<>"")', '"GENERAL expenses cannot carry a batch"'),
            (f'AND({b}<>"",COUNTIF({tr("tBatches", "Batch ID")},{b})=0)', '"Batch is not in BATCH_REGISTER"'),
            (f'AND({b}<>"",COUNTIFS({tr("tBatches", "Batch ID")},{b},{tr("tBatches", "Pen ID")},{p})=0)',
             f'"Batch "&{b}&" belongs to pen "&{bl("Pen ID", b)}'),
            (f'{me(T, "Payment Method")}=""', '"Missing payment method"'),
        ], '"OK"'))),
]
ws = WS["OTHER_EXPENSES"]
sheet_title(ws, "OTHER EXPENSES", "All costs other than feed & drugs (chicks, labour, power, repairs...). Use GENERAL "
            "for farm-wide costs; add a Batch ID when the cost belongs to one production cycle.")
build_table(ws, T, ex_cols, exp_rows)
check_cf(ws, col_letter(ex_cols, "Check"))

# --------------------------------------------------------------------------------------
# CASHBOOK
# --------------------------------------------------------------------------------------
T = "tCash"
direction = f'IFERROR(INDEX({tr("tCashTypes", "Direction")},MATCH({me(T, "Type")},{tr("tCashTypes", "Cash Type")},0)),"")'
cb_cols = [
    Col("Date", 11, DMY, dv=DV_DATE),
    Col("Entry ID", 11, formula=f'"CB-"&TEXT(ROW()-ROW({T}[[#Headers],[Date]]),"00000")'),
    Col("Type", 20, dv=dv_list("L_CashTypes")),
    Col("Category", 20, dv=dv_list("L_ExpCat", False)),
    Col("Description", 32),
    Col("Pen ID / General", 10, dv=dv_list("L_PensGen")),
    Col("Money In", 13, NGN, dv=DV_NUM),
    Col("Money Out", 13, NGN, dv=DV_NUM),
    Col("Running Balance", 14, NGN, formula=f"SUM(INDEX({tr(T, 'Money In')},1):{me(T, 'Money In')})"
        f"-SUM(INDEX({tr(T, 'Money Out')},1):{me(T, 'Money Out')})",
        note="Cumulative balance in row order - enter cash movements in date order."),
    Col("Payment Method", 10, dv=dv_list("L_Pay")),
    Col("Reference", 12),
    Col("Check", 28, formula=blank_guard(T, ["Date", "Type", "Money In", "Money Out"], nested_if(date_checks(T) + [
        (f'{me(T, "Type")}=""', '"Missing type"'),
        (f"AND(N({me(T, 'Money In')})>0,N({me(T, 'Money Out')})>0)", '"Use one row for In and another for Out"'),
        (f"N({me(T, 'Money In')})+N({me(T, 'Money Out')})=0", '"Enter an amount"'),
        (f'AND({direction}="In",N({me(T, "Money Out")})>0)', '"This type is money IN"'),
        (f'AND({direction}="Out",N({me(T, "Money In")})>0)', '"This type is money OUT"'),
        (f"{me(T, 'Running Balance')}<0", '"Balance is negative - check entries"'),
        (f"AND(ROW()>ROW({T}[[#Headers],[Date]])+1,{me(T, 'Date')}<INDEX({tr(T, 'Date')},ROW()-ROW({T}[[#Headers],[Date]])-1))",
         '"Date earlier than previous row"'),
    ], '"OK"'))),
]
ws = WS["CASHBOOK"]
sheet_title(ws, "CASHBOOK", "Actual money in and out (cash + bank) in date order. This is cash movement, NOT profit: "
            "buying feed is cash out today but becomes cost only when the feed is issued.")
build_table(ws, T, cb_cols, cash_rows)
check_cf(ws, col_letter(cb_cols, "Check"))

# --------------------------------------------------------------------------------------
# STOCK_MOVEMENTS
# --------------------------------------------------------------------------------------
T = "tMoves"
sm_cols = [
    Col("Date", 11, DMY, dv=DV_DATE),
    Col("Movement ID", 11, formula=f'"SM-"&TEXT(ROW()-ROW({T}[[#Headers],[Date]]),"00000")'),
    Col("Pen ID", 8, dv=dv_list("L_Pens")),
    Col("Batch ID", 18, dv=dv_list("L_Batches")),
    Col("Movement Type", 14, dv=dv_list("L_MoveTypes")),
    Col("Quantity In", 9, INT, dv=DV_INT),
    Col("Quantity Out", 9, INT, dv=DV_INT),
    Col("Related Batch ID", 18, dv=dv_list("L_Batches"), note="For transfers: the other batch involved."),
    Col("Notes", 34),
    Col("Check", 30, formula=blank_guard(T, ["Date", "Pen ID", "Batch ID", "Quantity In", "Quantity Out"],
        nested_if(date_checks(T) + pen_batch_checks(T) + [
            (f'{me(T, "Movement Type")}=""', '"Missing movement type"'),
            (f"AND(N({me(T, 'Quantity In')})>0,N({me(T, 'Quantity Out')})>0)", '"Use separate rows for In and Out"'),
            (f"N({me(T, 'Quantity In')})+N({me(T, 'Quantity Out')})=0", '"Enter a quantity"'),
            (f'AND({me(T, "Related Batch ID")}<>"",COUNTIF({tr("tBatches", "Batch ID")},{me(T, "Related Batch ID")})=0)',
             '"Related batch not found"'),
        ], '"OK"'))),
]
ws = WS["STOCK_MOVEMENTS"]
sheet_title(ws, "STOCK MOVEMENTS", "Head-count changes that are not deaths or sales: transfers between batches, extra "
            "stock added, count corrections. A transfer = one Out row + one In row.")
build_table(ws, T, sm_cols, move_rows)
check_cf(ws, col_letter(sm_cols, "Check"))


# --------------------------------------------------------------------------------------
# Report-sheet helpers
# --------------------------------------------------------------------------------------
def hdr_row(ws, r, c0, headers, widths=None, fill=FILL_IN_HDR):
    for j, h in enumerate(headers):
        c = ws.cell(row=r, column=c0 + j, value=h)
        c.font = F_HDR
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BOX
        if widths:
            ws.column_dimensions[get_column_letter(c0 + j)].width = widths[j]
    ws.row_dimensions[r].height = 30


def put(ws, ref, value, fmt=None, font=F_BASE, fill=None, border=BOX, align=None):
    c = ws[ref]
    c.value = value
    style_cell(c, font=font, fill=fill, fmt=fmt, border=border, align=align)
    return c


def protect(ws):
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def section(ws, ref, text):
    c = ws[ref]
    c.value = text
    c.font = F_SECTION


# Criteria (named) used by DASHBOARD / REPORTS / TRENDS -----------------------------
hs["F1"], hs["G1"] = "Dashboard filter logic", ""
filter_logic = [
    ("DStart", '=IF(DB_Period="Today",TODAY(),IF(DB_Period="This Week",TODAY()-WEEKDAY(TODAY(),2)+1,'
               'IF(DB_Period="This Month",DATE(YEAR(TODAY()),MONTH(TODAY()),1),IF(DB_Period="This Year",'
               'DATE(YEAR(TODAY()),1,1),IF(DB_Period="Custom",IF(ISNUMBER(DB_From),DB_From,1),1)))))'),
    ("DEnd", '=IF(DB_Period="Today",TODAY(),IF(DB_Period="This Week",TODAY()-WEEKDAY(TODAY(),2)+7,'
             'IF(DB_Period="This Month",EOMONTH(TODAY(),0),IF(DB_Period="This Year",DATE(YEAR(TODAY()),12,31),'
             'IF(DB_Period="Custom",IF(ISNUMBER(DB_To),DB_To,2958465),2958465)))))'),
    ("FPen", f'=IF(OR(DB_Pen="",DB_Pen="All Pens"),{ALL},DB_Pen)'),
    ("FBatch", f'=IF(OR(DB_Batch="",DB_Batch="All Batches"),{ALL},DB_Batch)'),
    ("FAnimal", f'=IF(OR(DB_Animal="",DB_Animal="All Types"),{ALL},DB_Animal)'),
    ("FilterText", '=IF(DB_Period="All Time","All time",TEXT(DStart,"dd/mm/yyyy")&" to "&TEXT(DEnd,"dd/mm/yyyy"))'
                   '&"  |  Pen: "&IF(DB_Pen="","All Pens",DB_Pen)&"  |  Batch: "&IF(DB_Batch="","All Batches",DB_Batch)'
                   '&"  |  Animal: "&IF(DB_Animal="","All Types",DB_Animal)'),
]
for i, (nm, f) in enumerate(filter_logic):
    r = 2 + i
    hs[f"F{r}"] = nm
    hs[f"G{r}"] = f
    define(nm, f"HELPER!$G${r}")
hs["G2"].number_format = DMY
hs["G3"].number_format = DMY

# Table meta: date / pen / batch / animal column names for filtering
META = {
    "tFeedIss": ("Date", "Pen ID", "Batch ID", "Animal Type"),
    "tDrugUse": ("Date", "Pen ID", "Batch ID", "Animal Type"),
    "tMort": ("Date", "Pen ID", "Batch ID", "Animal Type"),
    "tProd": ("Date", "Pen ID", "Batch ID", "Animal Type"),
    "tSales": ("Date", "Pen ID", "Batch ID", "Animal Type"),
    "tExp": ("Date", "Pen ID / General", "Batch ID (if applicable)", "Animal Type"),
}


def crit(t, start="DStart", end="DEnd", pen="FPen", batch="FBatch", animal="FAnimal"):
    d, p, b, a = META[t]
    parts = [f'{tr(t, d)},">="&{start}', f'{tr(t, d)},"<="&{end}']
    if pen:
        parts.append(f"{tr(t, p)},{pen}")
    if batch:
        parts.append(f"{tr(t, b)},{batch}")
    if animal:
        parts.append(f"{tr(t, a)},{animal}")
    return ",".join(parts)


def sumf(t, col, extra="", **kw):
    e = ("," + extra) if extra else ""
    return f"SUMIFS({tr(t, col)},{crit(t, **kw)}{e})"


def countf(t, extra="", **kw):
    e = ("," + extra) if extra else ""
    return f"COUNTIFS({crit(t, **kw)}{e})"


# --------------------------------------------------------------------------------------
# INVENTORY
# --------------------------------------------------------------------------------------
ws = WS["INVENTORY"]
sheet_title(ws, "INVENTORY (automatic)", "Store balances calculated from purchases minus issues/usage. Nothing to "
            "type here - reorder levels are set on LISTS. Costing: moving weighted average.", legend=False)
section(ws, "A4", "FEED STORE (kg)")
fh = ["Feed Type", "Purchased kg", "Purchased Cost", "Issued kg", "Issued Cost (to pens)", "Remaining kg",
      "Remaining Bags (approx.)", "Current Avg Cost per kg", "Inventory Value", "Reorder Level kg", "Status",
      "Last Purchase"]
hdr_row(ws, 5, 1, fh, [22, 12, 14, 12, 14, 12, 11, 12, 14, 11, 16, 12], fill=FILL_AUTO_HDR)
FEED_SLOTS = 20
r0 = 6
for n in range(1, FEED_SLOTS + 1):
    r = r0 + n - 1
    A = f"$A{r}"
    g = lambda expr: f'=IF({A}="","",{expr})'
    put(ws, f"A{r}", f'=IFERROR(INDEX({tr("tFeedTypes", "Feed Type")},{n})&"","")', font=F_BOLD)
    put(ws, f"B{r}", g(f"SUMIFS({tr('tFeedPur', 'Total Quantity in kg')},{tr('tFeedPur', 'Feed Type')},{A})"), DEC1)
    put(ws, f"C{r}", g(f"SUMIFS({tr('tFeedPur', 'Total Cost')},{tr('tFeedPur', 'Feed Type')},{A})"), NGN)
    put(ws, f"D{r}", g(f"SUMIFS({tr('tFeedIss', 'Quantity Issued in kg')},{tr('tFeedIss', 'Feed Type')},{A})"), DEC1)
    put(ws, f"E{r}", g(f"SUMIFS({tr('tFeedIss', 'Total Cost')},{tr('tFeedIss', 'Feed Type')},{A})"), NGN)
    put(ws, f"F{r}", g(f"B{r}-D{r}"), DEC1)
    put(ws, f"G{r}", g(f'IFERROR(F{r}/INDEX({tr("tFeedTypes", "Default Bag Size kg")},{n}),"")'), DEC1)
    put(ws, f"H{r}", g(latest_avg("tFeedPur", "Feed Type", A, "2958465", "Avg Cost per kg After Purchase")), NGN2)
    put(ws, f"I{r}", g(f"C{r}-E{r}"), NGN)
    put(ws, f"J{r}", g(f"N(INDEX({tr('tFeedTypes', 'Reorder Level kg')},{n}))"), DEC1)
    put(ws, f"K{r}", g(f'IF(AND(B{r}=0,D{r}=0),"Not stocked",IF(F{r}<-0.0001,"NEGATIVE - CHECK",IF(F{r}<=J{r},"REORDER","OK")))'), font=F_BOLD)
    put(ws, f"L{r}", g(f'IF(COUNTIF({tr("tFeedPur", "Feed Type")},{A})=0,"",_xlfn.MAXIFS({tr("tFeedPur", "Date")},'
                       f'{tr("tFeedPur", "Feed Type")},{A}))'), DMY)
rt = r0 + FEED_SLOTS
FEED_TOT = rt
put(ws, f"A{rt}", "TOTAL FEED", font=F_BOLD, fill=FILL_SECTION)
for L_, fmt in zip("BCDEFGHIJKL", [DEC1, NGN, DEC1, NGN, DEC1, None, None, NGN, None, None, None]):
    if L_ in "BCDEFI":
        put(ws, f"{L_}{rt}", f"=SUM({L_}{r0}:{L_}{rt - 1})", fmt, font=F_BOLD, fill=FILL_SECTION)
    else:
        put(ws, f"{L_}{rt}", None, fill=FILL_SECTION)
put(ws, f"K{rt}", f'=COUNTIF(K{r0}:K{rt - 1},"REORDER")&" to reorder"', font=F_BOLD, fill=FILL_SECTION)
FEED_STATUS_RNG = f"INVENTORY!$K${r0}:$K${rt - 1}"
for rng_ in (f"K{r0}:K{rt - 1}",):
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'K{r0}="REORDER"'], fill=FILL_AMBER,
                                                    font=Font(name=FONT, bold=True, color="9C5700")))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT(K{r0},8)="NEGATIVE"'], fill=FILL_RED,
                                                    font=Font(name=FONT, bold=True, color=RED_DARK)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'K{r0}="OK"'], font=Font(name=FONT, color="006100")))

d0 = rt + 3
section(ws, f"A{d0}", "DRUG STORE (in each drug's stock unit)")
dh = ["Drug Name", "Unit", "Purchased Qty", "Purchased Cost", "Used Qty", "Used Cost (to pens)", "Remaining Qty",
      "Current Avg Cost per Unit", "Inventory Value", "Reorder Level", "Stock Status", "Nearest Unexpired Expiry",
      "Expiry Alert"]
hdr_row(ws, d0 + 1, 1, dh, fill=FILL_AUTO_HDR)
ws.column_dimensions["M"].width = 22
DRUG_SLOTS = 50
r1 = d0 + 2
for n in range(1, DRUG_SLOTS + 1):
    r = r1 + n - 1
    A = f"$A{r}"
    g = lambda expr: f'=IF({A}="","",{expr})'
    put(ws, f"A{r}", f'=IFERROR(INDEX({tr("tDrugs", "Drug Name")},{n})&"","")', font=F_BOLD)
    put(ws, f"B{r}", g(f'INDEX({tr("tDrugs", "Stock Unit")},{n})&""'))
    put(ws, f"C{r}", g(f"SUMIFS({tr('tDrugPur', 'Quantity Purchased')},{tr('tDrugPur', 'Drug Name')},{A})"), DEC1)
    put(ws, f"D{r}", g(f"SUMIFS({tr('tDrugPur', 'Total Cost')},{tr('tDrugPur', 'Drug Name')},{A})"), NGN)
    put(ws, f"E{r}", g(f"SUMIFS({tr('tDrugUse', 'Quantity Used')},{tr('tDrugUse', 'Drug Name')},{A})"), DEC1)
    put(ws, f"F{r}", g(f"SUMIFS({tr('tDrugUse', 'Total Cost')},{tr('tDrugUse', 'Drug Name')},{A})"), NGN)
    put(ws, f"G{r}", g(f"C{r}-E{r}"), DEC1)
    put(ws, f"H{r}", g(latest_avg("tDrugPur", "Drug Name", A, "2958465", "Avg Cost per Unit After Purchase")), NGN2)
    put(ws, f"I{r}", g(f"D{r}-F{r}"), NGN)
    put(ws, f"J{r}", g(f"N(INDEX({tr('tDrugs', 'Reorder Level')},{n}))"), DEC1)
    put(ws, f"K{r}", g(f'IF(AND(C{r}=0,E{r}=0),"Not stocked",IF(G{r}<-0.0001,"NEGATIVE - CHECK",IF(G{r}<=J{r},"REORDER","OK")))'), font=F_BOLD)
    nearest = (f"_xlfn.MINIFS({tr('tDrugPur', 'Expiry Date')},{tr('tDrugPur', 'Drug Name')},{A},"
               f"{tr('tDrugPur', 'Expiry Date')},\">=\"&TODAY(),{tr('tDrugPur', 'Expiry Status')},\"<>Used up\")")
    put(ws, f"L{r}", g(f'IF({nearest}=0,"",{nearest})'), DMY)
    exp_cnt = f"COUNTIFS({tr('tDrugPur', 'Drug Name')},{A},{tr('tDrugPur', 'Expiry Status')},\"EXPIRED\")"
    soon_cnt = f"COUNTIFS({tr('tDrugPur', 'Drug Name')},{A},{tr('tDrugPur', 'Expiry Status')},\"Expires*\")"
    put(ws, f"M{r}", g(f'IF({exp_cnt}>0,{exp_cnt}&" EXPIRED lot(s) in stock",IF({soon_cnt}>0,"Expiring within "'
                       f'&ExpiryWarnDays&" days","OK"))'), font=F_BOLD)
rt2 = r1 + DRUG_SLOTS
DRUG_TOT = rt2
put(ws, f"A{rt2}", "TOTAL DRUGS", font=F_BOLD, fill=FILL_SECTION)
for L_ in "BCDEFGHIJKLM":
    put(ws, f"{L_}{rt2}", None, fill=FILL_SECTION)
for L_ in "DFI":
    put(ws, f"{L_}{rt2}", f"=SUM({L_}{r1}:{L_}{rt2 - 1})", NGN, font=F_BOLD, fill=FILL_SECTION)
put(ws, f"K{rt2}", f'=COUNTIF(K{r1}:K{rt2 - 1},"REORDER")&" to reorder"', font=F_BOLD, fill=FILL_SECTION)
DRUG_STATUS_RNG = f"INVENTORY!$K${r1}:$K${rt2 - 1}"
DRUG_ALERT_RNG = f"INVENTORY!$M${r1}:$M${rt2 - 1}"
rng_ = f"K{r1}:K{rt2 - 1}"
ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'K{r1}="REORDER"'], fill=FILL_AMBER,
                                                font=Font(name=FONT, bold=True, color="9C5700")))
ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT(K{r1},8)="NEGATIVE"'], fill=FILL_RED))
ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'K{r1}="OK"'], font=Font(name=FONT, color="006100")))
rng_ = f"M{r1}:M{rt2 - 1}"
ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'ISNUMBER(SEARCH("EXPIRED",M{r1}))'], fill=FILL_RED,
                                                font=Font(name=FONT, bold=True, color=RED_DARK)))
ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT(M{r1},8)="Expiring"'], fill=FILL_AMBER))

rn = rt2 + 2
put(ws, f"A{rn}", "TOTAL INVENTORY VALUE", font=F_BOLD, fill=FILL_SECTION)
put(ws, f"B{rn}", f"=I{FEED_TOT}+I{DRUG_TOT}", NGN, font=F_BOLD, fill=FILL_SECTION)
ws[f"A{rn + 2}"] = ("Costing method: moving (perpetual) weighted average. After each purchase: new average = (stock in "
                    "store x old average + purchase cost) / (stock in store + quantity bought). Every issue/usage is "
                    "valued at the average in force on its date (see the grey columns on FEED_PURCHASES / DRUG_PURCHASES).")
ws[f"A{rn + 3}"] = ("Inventory Value = Purchased Cost - Cost already charged to pens, so purchases always reconcile "
                    "exactly to (cost consumed + stock still in store) = Remaining x Current Avg Cost.")
for rr in (rn + 2, rn + 3):
    ws[f"A{rr}"].font = F_SUB
INV_TOTAL = f"INVENTORY!$B${rn}"
FEED_KG_TOTAL = f"INVENTORY!$F${FEED_TOT}"
FEED_VAL_TOTAL = f"INVENTORY!$I${FEED_TOT}"
DRUG_VAL_TOTAL = f"INVENTORY!$I${DRUG_TOT}"
INV_FEED_FIRST, INV_FEED_LAST = r0, FEED_TOT - 1
ws.freeze_panes = "B6"
protect(ws)

# --------------------------------------------------------------------------------------
# BATCH_PROFITABILITY
# --------------------------------------------------------------------------------------
ws = WS["BATCH_PROFITABILITY"]
sheet_title(ws, "BATCH PROFITABILITY (automatic, all-time per batch)",
            "Every batch ever recorded - active and closed - with its full-cycle costs and revenue. Direct costs only: "
            "general farm overheads are shown on the DASHBOARD, not allocated to batches.", legend=False)
bh = ["Batch ID", "Pen ID", "Animal Type", "Start Date", "Status", "Closing Date", "Age (days)", "Initial Qty",
      "Added / Transfers In", "Deaths", "Heads Sold", "Transfers Out", "Current Qty", "Mortality %", "Feed Used kg",
      "Feed Cost", "Drug Cost", "Other Direct Costs", "Total Cost", "Sales Revenue", "Profit / Loss",
      "Cost per Animal Placed", "Revenue per Animal Sold", "Feed per Animal Placed kg", "Latest Avg Weight kg"]
bw = [18, 8, 10, 11, 8, 11, 7, 9, 9, 8, 8, 9, 9, 9, 10, 13, 12, 13, 14, 14, 14, 11, 11, 10, 10]
HR = 4
hdr_row(ws, HR, 1, bh, bw, fill=FILL_AUTO_HDR)
BATCH_SLOTS = 300
rb0 = HR + 1
for n in range(1, BATCH_SLOTS + 1):
    r = rb0 + n - 1
    A = f"$A{r}"
    g = lambda expr: f'=IF({A}="","",{expr})'
    idx = lambda col: f"INDEX({tr('tBatches', col)},{n})"
    put(ws, f"A{r}", f'=IFERROR(INDEX({tr("tBatches", "Batch ID")},{n})&"","")', font=F_BOLD)
    put(ws, f"B{r}", g(f'{idx("Pen ID")}&""'))
    put(ws, f"C{r}", g(f'{idx("Animal Type")}&""'))
    put(ws, f"D{r}", g(idx("Batch Start Date")), DMY)
    put(ws, f"E{r}", g(f'{idx("Status")}&""'))
    put(ws, f"F{r}", g(f'IF(N({idx("Closing Date")})=0,"",{idx("Closing Date")})'), DMY)
    put(ws, f"G{r}", g(idx("Age (days)")), INT)
    put(ws, f"H{r}", g(f'N({idx("Initial Quantity")})'), INT)
    put(ws, f"I{r}", g(idx("Transfers In / Added")), INT)
    put(ws, f"J{r}", g(idx("Deaths")), INT)
    put(ws, f"K{r}", g(idx("Heads Sold")), INT)
    put(ws, f"L{r}", g(idx("Transfers Out")), INT)
    put(ws, f"M{r}", g(idx("Current Quantity")), INT)
    put(ws, f"N{r}", g(f"IFERROR(J{r}/(H{r}+I{r}),0)"), PCT)
    put(ws, f"O{r}", g(f"SUMIFS({tr('tFeedIss', 'Quantity Issued in kg')},{tr('tFeedIss', 'Batch ID')},{A})"), DEC1)
    put(ws, f"P{r}", g(f"SUMIFS({tr('tFeedIss', 'Total Cost')},{tr('tFeedIss', 'Batch ID')},{A})"), NGN)
    put(ws, f"Q{r}", g(f"SUMIFS({tr('tDrugUse', 'Total Cost')},{tr('tDrugUse', 'Batch ID')},{A})"), NGN)
    put(ws, f"R{r}", g(f"SUMIFS({tr('tExp', 'Amount')},{tr('tExp', 'Batch ID (if applicable)')},{A})"), NGN)
    put(ws, f"S{r}", g(f"P{r}+Q{r}+R{r}"), NGN, font=F_BOLD)
    put(ws, f"T{r}", g(f"SUMIFS({tr('tSales', 'Total Revenue')},{tr('tSales', 'Batch ID')},{A})"), NGN, font=F_BOLD)
    put(ws, f"U{r}", g(f"T{r}-S{r}"), NGN, font=F_BOLD)
    put(ws, f"V{r}", g(f'IFERROR(S{r}/(H{r}+I{r}),"")'), NGN)
    put(ws, f"W{r}", g(f'IF(K{r}=0,"",T{r}/K{r})'), NGN)
    put(ws, f"X{r}", g(f'IFERROR(O{r}/(H{r}+I{r}),"")'), DEC2)
    lastw = (f"_xlfn.MAXIFS({tr('tProd', 'Date')},{tr('tProd', 'Batch ID')},{A},"
             f"{tr('tProd', 'Measurement Type')},\"Weighing\")")
    put(ws, f"Y{r}", g(f'IF({lastw}=0,"",IFERROR(AVERAGEIFS({tr("tProd", "Average Weight kg")},{tr("tProd", "Batch ID")},'
                       f'{A},{tr("tProd", "Measurement Type")},"Weighing",{tr("tProd", "Date")},{lastw}),""))'), WT)
rbl = rb0 + BATCH_SLOTS - 1
put(ws, "A3", "TOTAL (all batches)", font=F_BOLD, fill=FILL_SECTION)
for j in range(2, 26):
    L_ = get_column_letter(j)
    fmt = ws[f"{L_}{rb0}"].number_format
    if L_ in "HIJKLMOPQRSTU":
        put(ws, f"{L_}3", f"=SUM({L_}{rb0}:{L_}{rbl})", fmt, font=F_BOLD, fill=FILL_SECTION)
    else:
        put(ws, f"{L_}3", None, fill=FILL_SECTION)
put(ws, "N3", "=IFERROR(J3/(H3+I3),0)", PCT, font=F_BOLD, fill=FILL_SECTION)
ws.conditional_formatting.add(f"U3:U{rbl}", CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED,
                                                        font=Font(name=FONT, bold=True, color=RED_DARK)))
ws.conditional_formatting.add(f"E{rb0}:E{rbl}", FormulaRule(formula=[f'E{rb0}="Active"'], fill=FILL_GREEN))
ws.auto_filter.ref = f"A{HR}:Y{rbl}"
ws.freeze_panes = f"B{rb0}"
protect(ws)

# --------------------------------------------------------------------------------------
# TRENDS (monthly + daily tables, charts)
# --------------------------------------------------------------------------------------
ws = WS["TRENDS"]
sheet_title(ws, "TRENDS", "Month-by-month and day-by-day totals. Follows the Pen / Batch / Animal filters chosen on "
            "the DASHBOARD. Change the first month in the yellow cell.", legend=False)
put(ws, "A4", "First month shown:", font=F_BOLD, border=None)
put(ws, "C4", "=DATE(YEAR(TODAY()),MONTH(TODAY())-11,1)", "mmm yyyy", font=F_INPUT, fill=FILL_INPUT)
ws["C4"].protection = Protection(locked=False)
define("TrendStart", "TRENDS!$C$4")
put(ws, "E4", "=FilterText", font=F_SUB, border=None)
th = ["Month", "From", "To", "Feed Used kg", "Feed Cost", "Drug Cost", "Other Expenses", "Total Cost", "Revenue",
      "Profit / Loss", "Deaths", "Heads Sold", "Output Qty"]
tw = [11, 11, 11, 11, 13, 12, 13, 14, 14, 14, 9, 9, 10]
hdr_row(ws, 6, 1, th, tw, fill=FILL_AUTO_HDR)
M0 = 7
for i in range(12):
    r = M0 + i
    put(ws, f"B{r}", f"=EDATE(TrendStart,{i})", DMY)
    put(ws, f"C{r}", f"=EOMONTH(B{r},0)", DMY)
    put(ws, f"A{r}", f'=TEXT(B{r},"mmm yyyy")', font=F_BOLD)
    kw = dict(start=f"$B{r}", end=f"$C{r}")
    put(ws, f"D{r}", "=" + sumf("tFeedIss", "Quantity Issued in kg", **kw), DEC1)
    put(ws, f"E{r}", "=" + sumf("tFeedIss", "Total Cost", **kw), NGN)
    put(ws, f"F{r}", "=" + sumf("tDrugUse", "Total Cost", **kw), NGN)
    put(ws, f"G{r}", "=" + sumf("tExp", "Amount", **kw), NGN)
    put(ws, f"H{r}", f"=E{r}+F{r}+G{r}", NGN)
    put(ws, f"I{r}", "=" + sumf("tSales", "Total Revenue", **kw), NGN)
    put(ws, f"J{r}", f"=I{r}-H{r}", NGN, font=F_BOLD)
    put(ws, f"K{r}", "=" + sumf("tMort", "Number of Deaths", **kw), INT)
    put(ws, f"L{r}", "=" + sumf("tSales", "Heads Removed", **kw), INT)
    put(ws, f"M{r}", "=" + sumf("tProd", "Quantity", extra=f'{tr("tProd", "Measurement Type")},"<>Weighing"', **kw), DEC1)
M1 = M0 + 11
put(ws, f"A{M1 + 1}", "TOTAL", font=F_BOLD, fill=FILL_SECTION)
for L_ in "BC":
    put(ws, f"{L_}{M1 + 1}", None, fill=FILL_SECTION)
for L_ in "DEFGHIJKLM":
    put(ws, f"{L_}{M1 + 1}", f"=SUM({L_}{M0}:{L_}{M1})", ws[f"{L_}{M0}"].number_format, font=F_BOLD, fill=FILL_SECTION)
ws.conditional_formatting.add(f"J{M0}:J{M1 + 1}", CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED))

DD0 = M1 + 4
section(ws, f"A{DD0 - 1}", "DAILY ACTIVITY - last 14 days up to the dashboard period end (or today)")
dh2 = ["Day", "Date", "", "Feed Used kg", "Feed Cost", "Drug Cost", "Other Expenses", "Total Cost", "Revenue",
       "Profit / Loss", "Deaths", "Heads Sold", "Output Qty"]
hdr_row(ws, DD0, 1, dh2, fill=FILL_AUTO_HDR)
for i in range(14):
    r = DD0 + 1 + i
    put(ws, f"B{r}", f"=MIN(DEnd,TODAY())-{13 - i}", DMY)
    put(ws, f"A{r}", f'=TEXT(B{r},"ddd dd mmm")', font=F_BOLD)
    put(ws, f"C{r}", None)
    kw = dict(start=f"$B{r}", end=f"$B{r}")
    put(ws, f"D{r}", "=" + sumf("tFeedIss", "Quantity Issued in kg", **kw), DEC1)
    put(ws, f"E{r}", "=" + sumf("tFeedIss", "Total Cost", **kw), NGN)
    put(ws, f"F{r}", "=" + sumf("tDrugUse", "Total Cost", **kw), NGN)
    put(ws, f"G{r}", "=" + sumf("tExp", "Amount", **kw), NGN)
    put(ws, f"H{r}", f"=E{r}+F{r}+G{r}", NGN)
    put(ws, f"I{r}", "=" + sumf("tSales", "Total Revenue", **kw), NGN)
    put(ws, f"J{r}", f"=I{r}-H{r}", NGN, font=F_BOLD)
    put(ws, f"K{r}", "=" + sumf("tMort", "Number of Deaths", **kw), INT)
    put(ws, f"L{r}", "=" + sumf("tSales", "Heads Removed", **kw), INT)
    put(ws, f"M{r}", "=" + sumf("tProd", "Quantity", extra=f'{tr("tProd", "Measurement Type")},"<>Weighing"', **kw), DEC1)
ws.freeze_panes = "B7"


def mk_chart(kind, title, ytitle, series_cols, cats_ref, src_ws, r_first, r_last, w=16, h=7.5):
    ch = BarChart() if kind == "bar" else LineChart()
    if kind == "bar":
        ch.type = "col"
        ch.grouping = "clustered"
    ch.title = title
    ch.y_axis.title = ytitle
    ch.height, ch.width = h, w
    for c_ in series_cols:
        ref = Reference(src_ws, min_col=c_, min_row=r_first - 1, max_row=r_last)
        ch.add_data(ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    if kind == "line":
        for s_ in ch.series:
            s_.smooth = False
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"
    return ch


# Note: header row for monthly table is 6, data 7..18
cats = Reference(ws, min_col=1, min_row=M0, max_row=M1)
ch_feed = mk_chart("line", "Feed consumption per month (kg)", "kg", [4], cats, ws, M0, M1)
ch_rev = mk_chart("bar", "Revenue vs total cost per month", "₦", [9, 8], cats, ws, M0, M1)
ch_mort = mk_chart("bar", "Mortality trend (deaths per month)", "Deaths", [11], cats, ws, M0, M1)
ch_sales = mk_chart("line", "Sales trend (₦ per month)", "₦", [9], cats, ws, M0, M1)
ws.add_chart(ch_mort, "O6")
ws.add_chart(ch_sales, "O22")
inv = WS["INVENTORY"]
ch_inv = BarChart()
ch_inv.type = "bar"
ch_inv.title = "Feed in store (kg) by feed type"
ch_inv.height, ch_inv.width = 7.5, 16
ch_inv.add_data(Reference(inv, min_col=6, min_row=5, max_row=INV_FEED_LAST), titles_from_data=True)
ch_inv.set_categories(Reference(inv, min_col=1, min_row=INV_FEED_FIRST, max_row=INV_FEED_LAST))
ch_inv.x_axis.delete = False
ch_inv.y_axis.delete = False
ch_inv.legend = None
ws.add_chart(ch_inv, "O38")
protect(ws)

# --------------------------------------------------------------------------------------
# DASHBOARD
# --------------------------------------------------------------------------------------
ws = WS["DASHBOARD"]
ws.sheet_view.showGridLines = False
ws["B1"] = '="FARM DASHBOARD - "&FarmName'
ws["B1"].font = Font(name=FONT, size=18, bold=True, color=GREEN)
ws["B2"] = '="Showing: "&FilterText&"      (updated "&TEXT(TODAY(),"dd/mm/yyyy")&")"'
ws["B2"].font = F_SUB
ws.column_dimensions["A"].width = 2
for j in range(2, 15):
    ws.column_dimensions[get_column_letter(j)].width = 15


def label(ws, r, c0, text, **kw):
    """Card label spanning two columns (c0, c0+1)."""
    L0, L1 = get_column_letter(c0), get_column_letter(c0 + 1)
    put(ws, f"{L1}{r}", None, fill=kw.get("fill"))
    c = put(ws, f"{L0}{r}", text, **kw)
    ws.merge_cells(f"{L0}{r}:{L1}{r}")
    return c

section(ws, "B4", "FILTERS  (yellow cells - choose from dropdowns)")
filters = [("Period", "This Month", "DB_Period", '"Today,This Week,This Month,This Year,All Time,Custom"'),
           ("Custom From (if Period = Custom)", None, "DB_From", "date"),
           ("Custom To (if Period = Custom)", None, "DB_To", "date"),
           ("Pen", "All Pens", "DB_Pen", "L_PenFilter"),
           ("Batch", "All Batches", "DB_Batch", "L_BatchFilter"),
           ("Animal Type", "All Types", "DB_Animal", "L_AnimalFilter")]
for i, (lab, val, nm, src) in enumerate(filters):
    r = 5 + i
    label(ws, r, 2, lab, font=F_BOLD)
    c = put(ws, f"D{r}", val, DMY if src == "date" else None, font=F_INPUT, fill=FILL_INPUT)
    c.protection = Protection(locked=False)
    define(nm, f"DASHBOARD!$D${r}")
    if src == "date":
        apply_dv(ws, DV_DATE, f"D{r}")
    else:
        apply_dv(ws, dv_list(src), f"D{r}")
label(ws, 5, 5, "Period start", font=F_BOLD)
put(ws, "G5", '=IF(DB_Period="All Time","All records",DStart)', DMY)
label(ws, 6, 5, "Period end", font=F_BOLD)
put(ws, "G6", '=IF(DB_Period="All Time","All records",DEnd)', DMY)
label(ws, 8, 5, "Data warnings (see CHECKS)", font=F_BOLD)
put(ws, "G8", "=CHECKS!$C$4", INT, font=F_KPI)
label(ws, 9, 5, "Items to reorder", font=F_BOLD)
put(ws, "G9", f'=COUNTIF({FEED_STATUS_RNG},"REORDER")+COUNTIF({DRUG_STATUS_RNG},"REORDER")', INT, font=F_KPI)
label(ws, 10, 5, "Drugs with expired / expiring lots", font=F_BOLD)
put(ws, "G10", f'=COUNTIF({DRUG_ALERT_RNG},"*EXPIRED*")+COUNTIF({DRUG_ALERT_RNG},"Expiring*")', INT, font=F_KPI)
for ref in ("G8", "G9", "G10"):
    ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_AMBER))
put(ws, "I5", "How to use", font=F_BOLD, border=None)
tips = ["1. Pick a Period (or Custom + dates).", "2. Pick a Pen / Batch / Animal type, or leave 'All'.",
        "3. Everything below recalculates automatically.", "Cards marked * are farm-wide (not pen-filtered)."]
for i, t_ in enumerate(tips):
    put(ws, f"I{6 + i}", t_, font=F_SUB, border=None)

tB = "tBatches"
act_crit = (f'{tr(tB, "Status")},"Active",{tr(tB, "Pen ID")},FPen,{tr(tB, "Batch ID")},FBatch,'
            f'{tr(tB, "Animal Type")},FAnimal')
placed_in_period = (f'SUMIFS({tr(tB, "Total Placed")},{tr(tB, "Batch Start Date")},"<="&DEnd,{tr(tB, "Closing Date")},'
                    f'">="&DStart,{tr(tB, "Pen ID")},FPen,{tr(tB, "Batch ID")},FBatch,{tr(tB, "Animal Type")},FAnimal)'
                    f'+SUMIFS({tr(tB, "Total Placed")},{tr(tB, "Batch Start Date")},"<="&DEnd,{tr(tB, "Closing Date")},'
                    f'"",{tr(tB, "Pen ID")},FPen,{tr(tB, "Batch ID")},FBatch,{tr(tB, "Animal Type")},FAnimal)')
last_wt_date = (f'_xlfn.MAXIFS({tr("tProd", "Date")},{crit("tProd")},{tr("tProd", "Measurement Type")},"Weighing")')

cards = {
    2: ("CURRENT STATUS (today)", [
        ("Total current stock (animals)", f"=SUMIFS({tr(tB, 'Current Quantity')},{act_crit})", INT),
        ("Active pens", f'=IF(DB_Pen="All Pens",COUNTIFS({tr("tPens", "Status")},"Active",{tr("tPens", "Animal Type")},'
                        f'FAnimal),COUNTIFS({tr("tPens", "Status")},"Active",{tr("tPens", "Pen ID")},FPen))', INT),
        ("Active batches", f"=COUNTIFS({act_crit})", INT),
        ("Feed in store (kg) *", f"={FEED_KG_TOTAL}", DEC1),
        ("Feed stock value *", f"={FEED_VAL_TOTAL}", NGN),
        ("Drug stock value *", f"={DRUG_VAL_TOTAL}", NGN),
        ("Total inventory value *", f"={INV_TOTAL}", NGN),
        ("Total pens registered", f"=COUNTA({tr('tPens', 'Pen ID')})", INT),
        ("Batches ever recorded (incl. closed)", f"=COUNTA({tr(tB, 'Batch ID')})", INT),
    ]),
    5: ("FINANCIAL SUMMARY (selected period)", [
        ("Feed purchased *", f'=SUMIFS({tr("tFeedPur", "Total Cost")},{tr("tFeedPur", "Date")},">="&DStart,'
                             f'{tr("tFeedPur", "Date")},"<="&DEnd)', NGN),
        ("Drugs purchased *", f'=SUMIFS({tr("tDrugPur", "Total Cost")},{tr("tDrugPur", "Date")},">="&DStart,'
                              f'{tr("tDrugPur", "Date")},"<="&DEnd)', NGN),
        ("Total inventory purchases *", "=G13+G14", NGN),
        ("Feed cost (consumed)", "=" + sumf("tFeedIss", "Total Cost"), NGN),
        ("Drug cost (used)", "=" + sumf("tDrugUse", "Total Cost"), NGN),
        ("Other expenses", "=" + sumf("tExp", "Amount"), NGN),
        ("TOTAL COST", "=G16+G17+G18", NGN),
        ("TOTAL SALES", "=" + sumf("tSales", "Total Revenue"), NGN),
        ("PROFIT / LOSS", "=G20-G19", NGN),
        ("Cash balance at period end *", f'=SUMIFS({tr("tCash", "Money In")},{tr("tCash", "Date")},"<="&DEnd)'
                                         f'-SUMIFS({tr("tCash", "Money Out")},{tr("tCash", "Date")},"<="&DEnd)', NGN),
    ]),
    8: ("PRODUCTION SUMMARY (selected period)", [
        ("Deaths", "=" + sumf("tMort", "Number of Deaths"), INT),
        ("Mortality rate (deaths / animals placed)", f"=IFERROR(J13/({placed_in_period}),0)", PCT),
        ("Feed consumed (kg)", "=" + sumf("tFeedIss", "Quantity Issued in kg"), DEC1),
        ("Drug administrations (records)", "=" + countf("tDrugUse"), INT),
        ("Latest average weight (kg)", f'=IF({last_wt_date}=0,"",IFERROR(AVERAGEIFS({tr("tProd", "Average Weight kg")},'
                                       f'{crit("tProd")},{tr("tProd", "Measurement Type")},"Weighing",{tr("tProd", "Date")},'
                                       f'{last_wt_date}),""))', WT),
        ("Output quantity (eggs etc.)", "=" + sumf("tProd", "Quantity", extra=f'{tr("tProd", "Measurement Type")},"<>Weighing"'),
         DEC1),
        ("Animals sold (heads)", "=" + sumf("tSales", "Heads Removed"), INT),
        ("Animals placed in batches active in period", f"={placed_in_period}", INT),
    ]),
}
for c0, (title, items) in cards.items():
    vL = get_column_letter(c0 + 2)
    label(ws, 12, c0, title, font=F_HDR, fill=FILL_IN_HDR)
    put(ws, f"{vL}12", None, fill=FILL_IN_HDR)
    for i, (lab, f, fmt) in enumerate(items):
        r = 13 + i
        big = lab.isupper()
        label(ws, r, c0, lab, font=F_BOLD if big else F_BASE, fill=FILL_SECTION if big else None)
        put(ws, f"{vL}{r}", f, fmt, font=F_KPI, fill=FILL_SECTION if big else None,
            align=Alignment(horizontal="right"))
ws.conditional_formatting.add("G21", CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED,
                                                font=Font(name=FONT, size=12, bold=True, color=RED_DARK)))
ws.conditional_formatting.add("G21", CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_GREEN))

# Mortality mini-card (K:L)
label(ws, 12, 11, "MORTALITY (pen/batch filters)", font=F_HDR, fill=FILL_IN_HDR)
put(ws, "M12", None, fill=FILL_IN_HDR)
periods = [("Today", "TODAY()", "TODAY()"),
           ("This week", "TODAY()-WEEKDAY(TODAY(),2)+1", "TODAY()-WEEKDAY(TODAY(),2)+7"),
           ("This month", "DATE(YEAR(TODAY()),MONTH(TODAY()),1)", "EOMONTH(TODAY(),0)"),
           ("This year", "DATE(YEAR(TODAY()),1,1)", "DATE(YEAR(TODAY()),12,31)"),
           ("All time", "1", "2958465")]
for i, (lab, s_, e_) in enumerate(periods):
    r = 13 + i
    label(ws, r, 11, lab)
    put(ws, f"M{r}", "=" + sumf("tMort", "Number of Deaths", start=s_, end=e_), INT, font=F_KPI,
        align=Alignment(horizontal="right"))
label(ws, 18, 11, "Current stock (active batches)")
put(ws, "M18", "=D13", INT, font=F_KPI, align=Alignment(horizontal="right"))
ws["B23"] = "* = farm-wide figure (not affected by pen / batch / animal filters). Profit = sales - (feed & drugs actually " \
            "consumed + other expenses). Inventory purchases are NOT cost until issued."
ws["B23"].font = F_SUB

# Charts (source: TRENDS)
tws = WS["TRENDS"]
cats = Reference(tws, min_col=1, min_row=M0, max_row=M1)
ws.add_chart(mk_chart("bar", "Revenue vs total cost (last 12 months)", "₦", [9, 8], cats, tws, M0, M1, w=17, h=7.5), "B25")
ws.add_chart(mk_chart("line", "Feed consumption (kg per month)", "kg", [4], cats, tws, M0, M1, w=17, h=7.5), "H25")

# Pen comparison table
PC = 42
section(ws, f"B{PC - 1}", "PEN SUMMARY / COMPARISON (selected period; batch & animal filters apply) - factual figures, no ranking")
ph = ["Pen ID", "Pen Name", "Current Batch", "Status", "Current Stock", "Feed Used kg", "Feed Cost", "Drug Cost",
      "Other Cost", "Total Cost", "Revenue", "Profit / Loss", "Deaths"]
hdr_row(ws, PC, 2, ph, fill=FILL_AUTO_HDR)
PEN_SLOTS = 60
for n in range(1, PEN_SLOTS + 1):
    r = PC + n
    A = f"$B{r}"
    g = lambda expr: f'=IF({A}="","",{expr})'
    ip = lambda col: f"INDEX({tr('tPens', col)},{n})"
    kw = dict(pen=A)
    put(ws, f"B{r}", f'=IFERROR(INDEX({tr("tPens", "Pen ID")},{n})&"","")', font=F_BOLD)
    put(ws, f"C{r}", g(f'{ip("Pen Name")}&""'))
    put(ws, f"D{r}", g(f'{ip("Current Batch ID")}&""'))
    put(ws, f"E{r}", g(f'{ip("Status")}&""'))
    put(ws, f"F{r}", g(ip("Current Stock")), INT)
    put(ws, f"G{r}", g(sumf("tFeedIss", "Quantity Issued in kg", **kw)), DEC1)
    put(ws, f"H{r}", g(sumf("tFeedIss", "Total Cost", **kw)), NGN)
    put(ws, f"I{r}", g(sumf("tDrugUse", "Total Cost", **kw)), NGN)
    put(ws, f"J{r}", g(sumf("tExp", "Amount", **kw)), NGN)
    put(ws, f"K{r}", g(f"H{r}+I{r}+J{r}"), NGN, font=F_BOLD)
    put(ws, f"L{r}", g(sumf("tSales", "Total Revenue", **kw)), NGN, font=F_BOLD)
    put(ws, f"M{r}", g(f"L{r}-K{r}"), NGN, font=F_BOLD)
    put(ws, f"N{r}", g(sumf("tMort", "Number of Deaths", **kw)), INT)
rg = PC + PEN_SLOTS + 1
put(ws, f"B{rg}", "GENERAL (farm-wide)", font=F_BOLD, fill=FILL_SECTION)
for L_ in "CDEFGHI":
    put(ws, f"{L_}{rg}", None, fill=FILL_SECTION)
put(ws, f"J{rg}", "=" + sumf("tExp", "Amount", pen='"GENERAL"'), NGN, fill=FILL_SECTION)
put(ws, f"K{rg}", f"=J{rg}", NGN, font=F_BOLD, fill=FILL_SECTION)
put(ws, f"L{rg}", "=" + sumf("tSales", "Total Revenue", pen='"GENERAL"'), NGN, font=F_BOLD, fill=FILL_SECTION)
put(ws, f"M{rg}", f"=L{rg}-K{rg}", NGN, font=F_BOLD, fill=FILL_SECTION)
put(ws, f"N{rg}", None, fill=FILL_SECTION)
rT = rg + 1
put(ws, f"B{rT}", "TOTAL", font=F_BOLD, fill=FILL_SECTION)
for L_ in "CDE":
    put(ws, f"{L_}{rT}", None, fill=FILL_SECTION)
for L_ in "FGHIJKLMN":
    put(ws, f"{L_}{rT}", f"=SUM({L_}{PC + 1}:{L_}{rg})", ws[f"{L_}{PC + 1}"].number_format, font=F_BOLD, fill=FILL_SECTION)
ws.conditional_formatting.add(f"M{PC + 1}:M{rT}", CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED,
                                                             font=Font(name=FONT, bold=True, color=RED_DARK)))
ws.freeze_panes = "A4"
protect(ws)

# --------------------------------------------------------------------------------------
# REPORTS
# --------------------------------------------------------------------------------------
ws = WS["REPORTS"]
sheet_title(ws, "REPORTS", "Breakdowns for the period and Pen / Batch / Animal filters chosen on the DASHBOARD. Monthly & "
            "daily views are on TRENDS; batch results on BATCH_PROFITABILITY; stock on INVENTORY.", legend=False)
put(ws, "A3", "=\"Filters in use: \"&FilterText", font=F_BOLD, border=None)
for k, v in {"A": 30, "B": 13, "C": 14, "D": 11, "E": 9, "F": 3, "G": 30, "H": 13, "I": 14, "J": 11, "K": 9}.items():
    ws.column_dimensions[k].width = v


def report_block(ws, r, c0, title, headers, list_tbl, list_col, slots, value_fns, fmts, other_fns=None):
    """Category breakdown: one row per list item, plus 'Not in list' and a TOTAL row.
    value_fns: functions(label_ref) -> formula (without '=')."""
    cL = get_column_letter(c0)
    section(ws, f"{cL}{r}", title)
    hdr_row(ws, r + 1, c0, headers, fill=FILL_AUTO_HDR)
    first = r + 2
    for n in range(1, slots + 1):
        rr = first + n - 1
        lab = f"${cL}{rr}"
        put(ws, f"{cL}{rr}", f'=IFERROR(INDEX({tr(list_tbl, list_col)},{n})&"","")', font=F_BOLD)
        for j, fn in enumerate(value_fns):
            L_ = get_column_letter(c0 + 1 + j)
            put(ws, f"{L_}{rr}", f'=IF({lab}="","",{fn(lab, first, first + slots - 1, rr)})', fmts[j])
    ro = first + slots
    put(ws, f"{cL}{ro}", "Not in list / blank", font=F_SUB)
    rt_ = ro + 1
    put(ws, f"{cL}{rt_}", "TOTAL", font=F_BOLD, fill=FILL_SECTION)
    for j, fn in enumerate(value_fns):
        L_ = get_column_letter(c0 + 1 + j)
        if other_fns and other_fns[j] is not None:
            put(ws, f"{L_}{rt_}", "=" + other_fns[j](rt_, first, ro), fmts[j], font=F_BOLD, fill=FILL_SECTION)
            put(ws, f"{L_}{ro}", f"={L_}{rt_}-SUM({L_}{first}:{L_}{ro - 1})", fmts[j])
        else:
            put(ws, f"{L_}{rt_}", None, fill=FILL_SECTION)
            put(ws, f"{L_}{ro}", None)
    return rt_


def share(col_letter_, total_row_holder):
    return lambda lab, f, l, rr: f"IFERROR({col_letter_}{rr}/{col_letter_}${total_row_holder[0]},0)"


row = 5
# ---- Block 1: feed by type | drug by drug
SL1 = 20
tot_feed = [row + 2 + SL1 + 1]
tot_drug = [row + 2 + SL1 + 1]
f_fe = "tFeedIss"
report_block(ws, row, 1, "FEED CONSUMPTION BY FEED TYPE", ["Feed Type", "kg Issued", "Cost", "Issues (records)", "% of kg"],
             "tFeedTypes", "Feed Type", SL1,
             [lambda lab, f, l, rr: sumf(f_fe, "Quantity Issued in kg", extra=f"{tr(f_fe, 'Feed Type')},{lab}"),
              lambda lab, f, l, rr: sumf(f_fe, "Total Cost", extra=f"{tr(f_fe, 'Feed Type')},{lab}"),
              lambda lab, f, l, rr: countf(f_fe, extra=f"{tr(f_fe, 'Feed Type')},{lab}"),
              share("B", tot_feed)],
             [DEC1, NGN, INT, PCT],
             [lambda rt_, f, o: sumf(f_fe, "Quantity Issued in kg"), lambda rt_, f, o: sumf(f_fe, "Total Cost"),
              lambda rt_, f, o: countf(f_fe), None])
f_du = "tDrugUse"
report_block(ws, row, 7, "DRUG USAGE BY DRUG", ["Drug", "Qty Used", "Cost", "Administrations", "Unit"],
             "tDrugs", "Drug Name", SL1,
             [lambda lab, f, l, rr: sumf(f_du, "Quantity Used", extra=f"{tr(f_du, 'Drug Name')},{lab}"),
              lambda lab, f, l, rr: sumf(f_du, "Total Cost", extra=f"{tr(f_du, 'Drug Name')},{lab}"),
              lambda lab, f, l, rr: countf(f_du, extra=f"{tr(f_du, 'Drug Name')},{lab}"),
              lambda lab, f, l, rr: f'IFERROR(INDEX({tr("tDrugs", "Stock Unit")},MATCH({lab},{tr("tDrugs", "Drug Name")},0))&"","")'],
             [DEC1, NGN, INT, None],
             [None, lambda rt_, f, o: sumf(f_du, "Total Cost"), lambda rt_, f, o: countf(f_du), None])
row = row + 2 + SL1 + 4
# ---- Block 2: mortality by cause | sales by product
SL2 = 15
tot_mort = [row + 2 + SL2 + 1]
f_mo = "tMort"
report_block(ws, row, 1, "MORTALITY BY CAUSE", ["Cause", "Deaths", "Records", "% of deaths", ""],
             "tCauses", "Mortality Cause", SL2,
             [lambda lab, f, l, rr: sumf(f_mo, "Number of Deaths", extra=f"{tr(f_mo, 'Suspected/Recorded Cause')},{lab}"),
              lambda lab, f, l, rr: countf(f_mo, extra=f"{tr(f_mo, 'Suspected/Recorded Cause')},{lab}"),
              share("B", tot_mort)],
             [INT, INT, PCT],
             [lambda rt_, f, o: sumf(f_mo, "Number of Deaths"), lambda rt_, f, o: countf(f_mo), None])
f_sa = "tSales"
report_block(ws, row, 7, "SALES BY PRODUCT", ["Product", "Qty Sold", "Revenue", "Sales (records)", "Unit"],
             "tProducts", "Product", SL2,
             [lambda lab, f, l, rr: sumf(f_sa, "Quantity Sold", extra=f"{tr(f_sa, 'Product')},{lab}"),
              lambda lab, f, l, rr: sumf(f_sa, "Total Revenue", extra=f"{tr(f_sa, 'Product')},{lab}"),
              lambda lab, f, l, rr: countf(f_sa, extra=f"{tr(f_sa, 'Product')},{lab}"),
              lambda lab, f, l, rr: f'IFERROR(INDEX({tr("tProducts", "Sale Unit")},MATCH({lab},{tr("tProducts", "Product")},0))&"","")'],
             [DEC1, NGN, INT, None],
             [None, lambda rt_, f, o: sumf(f_sa, "Total Revenue"), lambda rt_, f, o: countf(f_sa), None])
row = row + 2 + SL2 + 4
# ---- Block 3: expenses by category | sales by customer
SL3 = 25
f_ex = "tExp"
report_block(ws, row, 1, "EXPENSES BY CATEGORY", ["Category", "Total", "Batch direct", "Pen-level", "General"],
             "tExpCat", "Expense Category", SL3,
             [lambda lab, f, l, rr: sumf(f_ex, "Amount", extra=f"{tr(f_ex, 'Category')},{lab}"),
              lambda lab, f, l, rr: sumf(f_ex, "Amount", extra=f"{tr(f_ex, 'Category')},{lab},{tr(f_ex, 'Cost Level')},\"Batch direct\""),
              lambda lab, f, l, rr: sumf(f_ex, "Amount", extra=f"{tr(f_ex, 'Category')},{lab},{tr(f_ex, 'Cost Level')},\"Pen-level\""),
              lambda lab, f, l, rr: sumf(f_ex, "Amount", extra=f"{tr(f_ex, 'Category')},{lab},{tr(f_ex, 'Cost Level')},\"General overhead\"")],
             [NGN, NGN, NGN, NGN],
             [lambda rt_, f, o: sumf(f_ex, "Amount"),
              lambda rt_, f, o: sumf(f_ex, "Amount", extra=f"{tr(f_ex, 'Cost Level')},\"Batch direct\""),
              lambda rt_, f, o: sumf(f_ex, "Amount", extra=f"{tr(f_ex, 'Cost Level')},\"Pen-level\""),
              lambda rt_, f, o: sumf(f_ex, "Amount", extra=f"{tr(f_ex, 'Cost Level')},\"General overhead\"")])
report_block(ws, row, 7, "SALES BY CUSTOMER", ["Customer", "Revenue", "Sales (records)", "On credit", ""],
             "tCustomers", "Customer", SL3,
             [lambda lab, f, l, rr: sumf(f_sa, "Total Revenue", extra=f"{tr(f_sa, 'Customer')},{lab}"),
              lambda lab, f, l, rr: countf(f_sa, extra=f"{tr(f_sa, 'Customer')},{lab}"),
              lambda lab, f, l, rr: sumf(f_sa, "Total Revenue", extra=f"{tr(f_sa, 'Customer')},{lab},{tr(f_sa, 'Payment Method')},\"Credit\"")],
             [NGN, INT, NGN],
             [lambda rt_, f, o: sumf(f_sa, "Total Revenue"), lambda rt_, f, o: countf(f_sa),
              lambda rt_, f, o: sumf(f_sa, "Total Revenue", extra=f"{tr(f_sa, 'Payment Method')},\"Credit\"")])
row = row + 2 + SL3 + 3
ws[f"A{row}"] = ("Tip: every log is an Excel Table, so Insert > PivotTable on any log gives further ad-hoc views "
                 "(e.g. feed by pen by week). Use the filter arrows on log headers to drill into records.")
ws[f"A{row}"].font = F_SUB
ws.freeze_panes = "A4"
protect(ws)

# --------------------------------------------------------------------------------------
# CHECKS
# --------------------------------------------------------------------------------------
ws = WS["CHECKS"]
sheet_title(ws, "CHECKS - data quality & safeguards", "Anything other than 0 needs attention. Flagged rows are "
            "highlighted red in the Check column of each sheet (use the filter arrow on 'Check').", legend=False)
put(ws, "A4", "TOTAL ROWS NEEDING ATTENTION", font=F_BOLD, fill=FILL_SECTION)
put(ws, "B4", None, fill=FILL_SECTION)
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 50
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 16
ws.column_dimensions["E"].width = 16
hdr_row(ws, 6, 1, ["Sheet", "Check", "Rows flagged"], fill=FILL_AUTO_HDR)
check_list = [("PEN_REGISTER", "tPens"), ("BATCH_REGISTER", "tBatches"), ("FEED_PURCHASES", "tFeedPur"),
              ("FEED_ISSUES", "tFeedIss"), ("DRUG_PURCHASES", "tDrugPur"), ("DRUG_USAGE", "tDrugUse"),
              ("MORTALITY_LOG", "tMort"), ("PRODUCTION_LOG", "tProd"), ("SALES", "tSales"),
              ("OTHER_EXPENSES", "tExp"), ("CASHBOOK", "tCash"), ("STOCK_MOVEMENTS", "tMoves")]
r = 7
for sh, t in check_list:
    put(ws, f"A{r}", sh, font=F_BOLD)
    put(ws, f"B{r}", "Rows whose Check column is not OK (missing IDs, wrong pen/batch, stock exceeded ...)")
    put(ws, f"C{r}", f'=COUNTIF({tr(t, "Check")},"?*")-COUNTIF({tr(t, "Check")},"OK")', INT, font=F_KPI)
    r += 1
extra = [
    ("INVENTORY", "Feed types with NEGATIVE stock", f'=COUNTIF({FEED_STATUS_RNG},"NEGATIVE*")'),
    ("INVENTORY", "Drugs with NEGATIVE stock", f'=COUNTIF({DRUG_STATUS_RNG},"NEGATIVE*")'),
    ("FEED", "Feed purchased/issued under a type missing from LISTS",
     f"=COUNTA({tr('tFeedPur', 'Feed Type')})+COUNTA({tr('tFeedIss', 'Feed Type')})"
     f"-SUMPRODUCT(COUNTIF({tr('tFeedPur', 'Feed Type')},{tr('tFeedTypes', 'Feed Type')}))"
     f"-SUMPRODUCT(COUNTIF({tr('tFeedIss', 'Feed Type')},{tr('tFeedTypes', 'Feed Type')}))"),
    ("DRUG_PURCHASES", "Lots EXPIRED with stock remaining", f'=COUNTIF({tr("tDrugPur", "Expiry Status")},"EXPIRED")'),
    ("DRUG_PURCHASES", '="Lots expiring within "&ExpiryWarnDays&" days"', f'=COUNTIF({tr("tDrugPur", "Expiry Status")},"Expires*")'),
    ("CASHBOOK", "Rows with a negative running balance", f'=COUNTIF({tr("tCash", "Running Balance")},"<0")'),
]
first_extra = r
for sh, lab, f in extra:
    put(ws, f"A{r}", sh, font=F_BOLD)
    put(ws, f"B{r}", lab)
    put(ws, f"C{r}", f, INT, font=F_KPI)
    r += 1
put(ws, "C4", f"=SUM(C7:C{first_extra + 2})", INT, font=F_KPI, fill=FILL_SECTION)
ws.conditional_formatting.add(f"C4:C{r - 1}", CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_RED,
                                                          font=Font(name=FONT, size=12, bold=True, color=RED_DARK)))
ws.conditional_formatting.add(f"C7:C{r - 1}", CellIsRule(operator="equal", formula=["0"], fill=FILL_GREEN))

r += 2
section(ws, f"A{r}", "CASH RECONCILIATION (all time) - logs vs cashbook")
r += 1
hdr_row(ws, r, 1, ["Item", "Explanation", "Per logs", "Per cashbook", "Difference"], fill=FILL_AUTO_HDR)
r += 1
rec = [
    ("Feed purchases paid", "FEED_PURCHASES not on Credit vs cashbook 'Feed Purchase'",
     f'SUMIFS({tr("tFeedPur", "Total Cost")},{tr("tFeedPur", "Payment Method")},"<>Credit")',
     f'SUMIFS({tr("tCash", "Money Out")},{tr("tCash", "Type")},"Feed Purchase")'),
    ("Drug purchases paid", "DRUG_PURCHASES not on Credit vs cashbook 'Drug Purchase'",
     f'SUMIFS({tr("tDrugPur", "Total Cost")},{tr("tDrugPur", "Payment Method")},"<>Credit")',
     f'SUMIFS({tr("tCash", "Money Out")},{tr("tCash", "Type")},"Drug Purchase")'),
    ("Other expenses paid", "OTHER_EXPENSES not on Credit vs cashbook 'Expense Payment'",
     f'SUMIFS({tr("tExp", "Amount")},{tr("tExp", "Payment Method")},"<>Credit")',
     f'SUMIFS({tr("tCash", "Money Out")},{tr("tCash", "Type")},"Expense Payment")'),
    ("Sales received", "SALES not on Credit vs cashbook 'Sale Receipt'",
     f'SUMIFS({tr("tSales", "Total Revenue")},{tr("tSales", "Payment Method")},"<>Credit")',
     f'SUMIFS({tr("tCash", "Money In")},{tr("tCash", "Type")},"Sale Receipt")'),
    ("Owed to suppliers (credit)", "Credit purchases & expenses minus 'Supplier Credit Payment'",
     f'SUMIFS({tr("tFeedPur", "Total Cost")},{tr("tFeedPur", "Payment Method")},"Credit")'
     f'+SUMIFS({tr("tDrugPur", "Total Cost")},{tr("tDrugPur", "Payment Method")},"Credit")'
     f'+SUMIFS({tr("tExp", "Amount")},{tr("tExp", "Payment Method")},"Credit")',
     f'SUMIFS({tr("tCash", "Money Out")},{tr("tCash", "Type")},"Supplier Credit Payment")'),
    ("Owed by customers (credit)", "Credit sales minus 'Customer Credit Receipt'",
     f'SUMIFS({tr("tSales", "Total Revenue")},{tr("tSales", "Payment Method")},"Credit")',
     f'SUMIFS({tr("tCash", "Money In")},{tr("tCash", "Type")},"Customer Credit Receipt")'),
]
for i, (lab, ex, a, b_) in enumerate(rec):
    put(ws, f"A{r}", lab, font=F_BOLD)
    put(ws, f"B{r}", ex, font=F_SUB)
    put(ws, f"C{r}", "=" + a, NGN)
    put(ws, f"D{r}", "=" + b_, NGN)
    put(ws, f"E{r}", f"=C{r}-D{r}", NGN, font=F_BOLD)
    if i < 4:
        ws.conditional_formatting.add(f"E{r}", FormulaRule(formula=[f"ABS(E{r})>0.5"], fill=FILL_AMBER))
    r += 1
ws[f"A{r + 1}"] = ("Differences in the first four lines mean a payment was recorded in one place but not the other. The "
                   "last two lines are outstanding credit balances (Difference = amount still owed).")
ws[f"A{r + 1}"].font = F_SUB
protect(ws)

# --------------------------------------------------------------------------------------
# README
# --------------------------------------------------------------------------------------
ws = WS["README"]
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 28
ws.column_dimensions["C"].width = 110
ws["B1"] = "FARM ACCOUNTING & INVENTORY MANAGEMENT SYSTEM"
ws["B1"].font = F_TITLE
ws["B2"] = "Enter transactions once. Everything else is calculated automatically."
ws["B2"].font = Font(name=FONT, size=11, italic=True, bold=True, color=GREY)

README = [
    ("H", "WHAT THIS WORKBOOK DOES"),
    ("", "Flow", "Purchase  ->  Inventory (store)  ->  Issue / Usage  ->  Pen & Batch  ->  Production  ->  Sale  ->  Profitability"),
    ("", "Hierarchy", "Farm -> Pen (permanent physical house, e.g. P001) -> Batch (one production cycle in a pen, e.g. "
                      "BROILER-SEP-001) -> Transactions. A pen keeps its ID for life; batches come and go."),
    ("", "Key idea", "Purchases are NOT charged to a pen. They go into the store. Cost reaches a pen/batch only when you "
                     "record an ISSUE (feed) or USAGE (drug). One purchase can therefore be shared by any number of pens."),
    ("", "No daily rows", "Record only what happened. If nothing happened in a pen today, enter nothing for it."),
    ("H", "COLOURS & CELLS"),
    ("", "Green header", "Input column - type or pick from the dropdown."),
    ("", "Grey header / grey cell", "Automatic formula - do not type. New rows fill these in by themselves."),
    ("", "Yellow cell", "A setting or filter you may change (DASHBOARD filters, LISTS settings, TRENDS first month)."),
    ("", "Check column", "Every log ends with a Check column: OK or a plain-English warning (red). CHECKS sheet counts them."),
    ("", "Protected sheets", "DASHBOARD, INVENTORY, BATCH_PROFITABILITY, REPORTS, TRENDS and CHECKS are protected "
                             "(no password) so formulas are not overwritten. Review > Unprotect Sheet if you must edit."),
    ("H", "SHEETS"),
    ("", "DASHBOARD", "Main page. Farm status, financial & production summary, mortality, charts and pen comparison. "
                      "Filter by period (Today / This Week / This Month / This Year / All Time / Custom), pen, batch, animal."),
    ("", "PEN_REGISTER", "List of pens. You enter ID, name, location, animal type. Current batch, stock and status are automatic."),
    ("", "BATCH_REGISTER", "One row per batch/cycle. You enter ID, pen, start date, initial quantity, status (& closing date). "
                           "Current quantity = initial + transfers in - deaths - animals sold - transfers out."),
    ("", "FEED_PURCHASES", "Every feed purchase (bags x bag size = kg; bags x price = cost)."),
    ("", "FEED_ISSUES", "Feed taken from the store to a pen/batch. Cost per kg is automatic (weighted average)."),
    ("", "DRUG_PURCHASES", "Every drug/vaccine purchase, with expiry date and lot number. Expiry warnings are automatic."),
    ("", "DRUG_USAGE", "Each administration per pen/batch with dosage and quantity. Cost is automatic."),
    ("", "MORTALITY_LOG", "Deaths - only on days they happen. Record the observed/recorded cause (use Unknown if unsure)."),
    ("", "PRODUCTION_LOG", "Weighings (average weight, number sampled) and output such as egg collection."),
    ("", "SALES", "Sales linked to pen & batch (or GENERAL for farm-wide items like manure)."),
    ("", "OTHER_EXPENSES", "Chicks/stock purchase, labour, electricity, repairs ... Pen ID or GENERAL; add Batch ID to charge a batch."),
    ("", "CASHBOOK", "Actual money in/out with running balance. Cash is not the same as profit (see below)."),
    ("", "INVENTORY", "Automatic store balances for feed and drugs, value, reorder and expiry status."),
    ("", "BATCH_PROFITABILITY", "Automatic full-cycle results for every batch ever recorded, active or closed."),
    ("", "STOCK_MOVEMENTS", "Transfers between batches, extra stock added, count corrections."),
    ("", "REPORTS", "Feed by type, drugs by drug, mortality by cause, sales by product & customer, expenses by category."),
    ("", "TRENDS", "Monthly and daily totals with charts (feed, revenue vs cost, mortality, sales, stock levels)."),
    ("", "CHECKS", "Data-quality counters and cash reconciliation. Aim for all zeros."),
    ("", "LISTS", "Master lists for all dropdowns (feed types + reorder levels, drugs + units, products, categories ...) and settings."),
    ("H", "HOW TO ..."),
    ("", "Add a new pen", "PEN_REGISTER: type the new Pen ID (e.g. P007) in the first empty row under the table, then name, "
                          "location, animal type. Never reuse or rename an ID that already has records."),
    ("", "Start a new batch", "BATCH_REGISTER: new row with a unique Batch ID (e.g. BROILER-OCT-001), Pen ID, animal type, start "
                              "date, initial quantity, Status = Active. Record the chick cost on OTHER_EXPENSES with the Batch ID."),
    ("", "Close a batch", "Set Status = Closed and enter the Closing Date. Do NOT delete anything - the batch's history stays "
                          "in every report and the pen shows Empty, ready for the next batch."),
    ("", "Record a purchase", "FEED_PURCHASES (feed type, bags, bag size, price per bag) or DRUG_PURCHASES (drug, quantity in "
                              "its stock unit, unit cost, expiry, lot no.). Also enter the payment in CASHBOOK unless bought on Credit."),
    ("", "Issue feed", "FEED_ISSUES: date, feed type, pen, batch, kg. Example: 500 kg bought once, then P001 40 kg, P002 60 kg, "
                       "P003 35 kg, P004 50 kg = four issue rows. If an issue exceeds the store, the Check column warns you."),
    ("", "Record drug usage", "DRUG_USAGE: date, drug, pen, batch, reason, dosage text, total quantity used. Same drug in 3 pens = 3 rows."),
    ("", "Record mortality", "MORTALITY_LOG: date, pen, batch, number of deaths, cause. Stock updates automatically."),
    ("", "Record a sale", "SALES: date, pen, batch, product, quantity, unit price, customer, payment method. Live-animal "
                          "products reduce batch stock automatically (set per product on LISTS). Add receipt to CASHBOOK."),
    ("", "Record an expense", "OTHER_EXPENSES: category, description, Pen ID or GENERAL, Batch ID if it belongs to one batch, amount."),
    ("", "Add a dropdown item", "LISTS: type the new item directly under the relevant list - dropdowns grow automatically."),
    ("H", "READING THE DASHBOARD"),
    ("", "Filters", "Choose Period, Pen, Batch and Animal Type in the yellow cells. All cards, the pen comparison, TRENDS and REPORTS follow them."),
    ("", "Current status", "Live stock, active pens/batches and store value as of today."),
    ("", "Financial", "Purchases (cash spent on stock) are shown separately from COST (what was consumed). "
                      "Profit/Loss = Sales - (feed issued + drugs used + other expenses) for the filter."),
    ("", "Pen comparison", "Factual figures per pen for the period. No ranking - interpret them yourself. GENERAL = farm-wide costs."),
    ("H", "COSTING & ACCOUNTING RULES"),
    ("", "Costing method", "Moving weighted average (fixed - never changes silently). Each purchase re-averages the stock "
                           "in store: (stock on hand x old average + purchase cost) / (stock on hand + quantity bought). "
                           "Example: 100 kg in store @ ₦500 and 200 kg bought @ ₦600 -> (100x500 + 200x600) / 300 = ₦566.67 per kg. "
                           "Every issue is costed at the average in force on its date; later purchases never change earlier issues."),
    ("", "Purchase order", "Keep FEED_PURCHASES and DRUG_PURCHASES in date order (the average is built row by row). A late-entered "
                           "old invoice must be inserted in its date position (right-click > Insert > Table Rows Above); "
                           "the Check column flags rows out of order."),
    ("", "Inventory value", "Purchased cost - cost already charged to pens. Purchases therefore always equal consumption + stock in store."),
    ("", "Purchase vs consumption", "Buying ₦500,000 of feed is NOT a ₦500,000 batch cost. It becomes cost only as it is issued. Unused drugs stay in inventory."),
    ("", "Cash vs profit", "CASHBOOK shows money actually received/paid (incl. credit settlements, capital, loans). Profit uses consumption. "
                           "CHECKS reconciles the cashbook with the logs."),
    ("", "Overheads", "GENERAL expenses (e.g. farm electricity) are in farm totals, not allocated to batches. Pen-level costs "
                      "without a batch are in pen totals only. Add a Batch ID to include a cost in batch profitability."),
    ("", "Units", "Feed is always stored in kg. Each drug has one Stock Unit (LISTS) - enter purchase and usage quantities in that unit."),
    ("H", "HISTORY, SCALE & GOOD PRACTICE"),
    ("", "History", "Never delete or overwrite old records when a batch closes. Closed batches stay in BATCH_PROFITABILITY for comparison."),
    ("", "Growth", "Logs are Excel Tables: they grow automatically and every formula uses whole-table references, so 5 or 50 pens and "
                   "tens of thousands of rows need no formula changes. Report sheets show up to 60 pens, 300 batches, "
                   "20 feed types and 50 drugs (extend with the generator script if ever needed)."),
    ("", "IDs", "Record IDs (FP-00001 ...) are numbered by row. Do not sort the logs - use the filter arrows to view subsets."),
    ("", "Dates", "Enter dates as DD/MM/YYYY. Invalid dates, text in number fields and negative amounts are rejected."),
    ("H", "SAMPLE DATA"),
    ("", "What it is", "All example rows are marked 'SAMPLE DATA' in Notes (pens P001-P006, six batches incl. one closed batch in P006, "
                       "a 20-bag x 25 kg Starter purchase shared across P001-P004, one drug used in three pens at different dosages)."),
    ("", "Remove it", "Delete the sample rows in each log (select rows inside the table > right-click > Delete > Table Rows), then "
                      "PEN_REGISTER / BATCH_REGISTER rows, then edit LISTS. Keep at least the table headers."),
]
r = 4
for item in README:
    if item[0] == "H":
        r += 1
        c = ws.cell(row=r, column=2, value=item[1])
        c.font = F_HDR
        c.fill = FILL_IN_HDR
        ws.cell(row=r, column=3).fill = FILL_IN_HDR
    else:
        a = ws.cell(row=r, column=2, value=item[1])
        a.font = F_BOLD
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=r, column=3, value=item[2])
        b.font = F_BASE
        b.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * max(1, -(-len(item[2]) // 120))
    r += 1

# base font for anything left unstyled + final touches
for w in wb.worksheets:
    w.sheet_view.zoomScale = 90
wb.active = wb.sheetnames.index("DASHBOARD")
for w in wb.worksheets:
    w.sheet_view.tabSelected = (w.title == "DASHBOARD")
wb.save(OUT)
print("saved", OUT)
